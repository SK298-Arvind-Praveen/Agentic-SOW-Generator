import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from app.agents.objective_agent import ObjectiveAgent
from app.agents.poc_writer_agent import POCWriterAgent, TemplateSection
from app.core.nodes import analyze_objective_node
from app.document.doc_reader import extract_supporting_documents, read_document


class SupportingDocumentFlowTests(unittest.TestCase):
    def _build_brd(self, root: str) -> Path:
        path = Path(root) / "specific_brd.docx"
        document = Document()
        document.add_heading("Settlement Exception Management", level=1)
        document.add_paragraph(
            "Operations analysts must triage failed settlement instructions from the Core Banking System."
        )
        table = document.add_table(rows=1, cols=4)
        for cell, value in zip(
            table.rows[0].cells,
            ("ID", "Module", "Deliverable", "Acceptance Evidence"),
        ):
            cell.text = value
        values = (
            "BRD-17",
            "Exception Triage",
            "Configurable ageing queue with maker-checker assignment",
            "Demonstrate routing for unmatched, partially matched, and rejected instructions",
        )
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
        nested = cells[2].add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Export the queue to CSV with the BRD-17 identifier"
        document.save(path)
        return path

    def _build_brd_bytes(self) -> io.BytesIO:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._build_brd(temp_dir)
            stream = io.BytesIO(path.read_bytes())
        stream.seek(0)
        return stream

    def test_docx_reader_preserves_specific_table_semantics_and_nested_content(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._build_brd(temp_dir)
            content = read_document(str(path))

        self.assertIn("HEADING: Settlement Exception Management", content)
        self.assertIn("COLUMNS: ID | Module | Deliverable | Acceptance Evidence", content)
        self.assertIn("ID: BRD-17", content)
        self.assertIn("Module: Exception Triage", content)
        self.assertIn("maker-checker assignment", content)
        self.assertIn("unmatched, partially matched, and rejected instructions", content)
        self.assertIn("Export the queue to CSV with the BRD-17 identifier", content)

    def test_legacy_doc_is_converted_before_docx_extraction(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            legacy = Path(temp_dir) / "requirements.doc"
            legacy.write_bytes(b"legacy-binary-placeholder")

            def fake_conversion(command, **_kwargs):
                out_dir = Path(command[command.index("--outdir") + 1])
                converted = Document()
                converted.add_paragraph("BRD-42 requires Cisco call tagging and ticket migration.")
                converted.save(out_dir / "requirements.docx")
                return SimpleNamespace(returncode=0, stdout="converted", stderr="")

            with (
                patch("app.document.doc_reader._find_soffice_for_legacy_doc", return_value="soffice"),
                patch("app.document.doc_reader.subprocess.run", side_effect=fake_conversion),
            ):
                content = read_document(str(legacy))

        self.assertIn("BRD-42", content)
        self.assertIn("Cisco call tagging", content)

    def test_legacy_doc_uses_antiword_when_libreoffice_is_unavailable(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            legacy = Path(temp_dir) / "requirements.doc"
            legacy.write_bytes(b"legacy-binary-placeholder")
            result = SimpleNamespace(
                returncode=0,
                stdout=b"Axis CRM migration, ticket classification and email workflow requirements.",
                stderr=b"",
            )
            with (
                patch("app.document.doc_reader._find_soffice_for_legacy_doc", return_value=None),
                patch("app.document.doc_reader.shutil.which", return_value="antiword"),
                patch("app.document.doc_reader.subprocess.run", return_value=result),
            ):
                content = read_document(str(legacy))

        self.assertIn("Axis CRM migration", content)
        self.assertIn("email workflow", content)

    def test_supporting_document_consolidation_retains_source_name_and_details(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._build_brd(temp_dir)
            content = extract_supporting_documents([str(path)])

        self.assertIn("SUPPORTING DOCUMENT 1: specific_brd.docx", content)
        self.assertIn("BRD-17", content)
        self.assertIn("Exception Triage", content)

    def test_xlsx_reader_preserves_sheet_rows_and_requirement_details(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "deliverables.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Scope"
            sheet.append(["Requirement ID", "Deliverable", "Acceptance Evidence"])
            sheet.append(["BRD-88", "Maker-checker exception queue", "Route rejected settlements"])
            workbook.save(path)
            workbook.close()

            content = read_document(str(path))

        self.assertIn("SHEET: Scope", content)
        self.assertIn("BRD-88", content)
        self.assertIn("Maker-checker exception queue", content)
        self.assertIn("Route rejected settlements", content)

    def test_document_only_analysis_prompt_makes_uploaded_corpus_authoritative(self):
        agent = ObjectiveAgent.__new__(ObjectiveAgent)
        prompt = agent._build_analysis_prompt(
            "",
            "BRD-17 requires a configurable maker-checker exception queue.",
        )
        self.assertIn("SUPPORTING DOCUMENT CORPUS (primary factual and scope authority)", prompt)
        self.assertIn("BRD-17", prompt)
        self.assertIn("No additional generation guidance supplied", prompt)
        self.assertNotIn("Build a comprehensive AWS-based solution", prompt)

    def test_uploaded_document_remains_primary_when_guidance_changes_scope_treatment(self):
        agent = ObjectiveAgent.__new__(ObjectiveAgent)
        prompt = agent._build_analysis_prompt(
            "Keep voice capability as future scope and prioritise workflow automation.",
            "BRD-21 lists voice capability and workflow automation as candidate deliverables.",
        )
        self.assertLess(
            prompt.index("SUPPORTING DOCUMENT CORPUS"),
            prompt.index("USER GENERATION GUIDANCE"),
        )
        self.assertIn("moving a document capability to future scope", prompt)
        self.assertIn("Keep voice capability as future scope", prompt)

    def test_objectives_subheading_is_removed_even_after_introductory_text(self):
        writer = POCWriterAgent.__new__(POCWriterAgent)
        content = writer._clean_content(
            "The engagement automates settlement exception handling.\n\n"
            "### Objectives\n\n- Route unmatched instructions.\n- Preserve BRD-17 evidence.",
            "3. Objective",
        )
        self.assertNotIn("### Objectives", content)
        self.assertIn("Route unmatched instructions", content)

    def test_writer_failure_fallback_keeps_extracted_deliverables(self):
        section = TemplateSection(
            "Deliverables",
            "Generate a deliverables table.",
            {"type": "GENERATED"},
            1,
        )
        content = POCWriterAgent._deterministic_fallback(section, {
            "project_overview": "Settlement exception management.",
            "key_deliverables": [
                "BRD-17 configurable ageing queue with maker-checker assignment"
            ],
        })
        self.assertIn("BRD-17", content)
        self.assertIn("maker-checker", content)
        self.assertIn("Module/Workstream", content)

    def test_analysis_node_does_not_inject_generic_objective_for_document_only_request(self):
        captured = {}

        class FakeObjectiveAgent:
            def __init__(self, _config):
                pass

            def analyze_objective(self, objective, supporting_context):
                captured["objective"] = objective
                captured["supporting_context"] = supporting_context
                return {"ui_required": False, "key_deliverables": ["BRD-17 exception queue"]}

            evaluate_input_quality = staticmethod(lambda *_args: {
                "accepted": True, "confidence": 0.9, "reason": "",
                "evidence_categories": {"key_deliverables": 1},
            })

        state = {
            "mode": "POC",
            "objective": "",
            "supporting_context": "BRD-17 exception queue",
            "metadata": {},
        }
        with patch("app.core.nodes.ObjectiveAgent", FakeObjectiveAgent):
            result = analyze_objective_node(state)

        self.assertEqual(captured["objective"], "")
        self.assertEqual(captured["supporting_context"], "BRD-17 exception queue")
        self.assertEqual(result["analyzed_requirements"]["key_deliverables"], ["BRD-17 exception queue"])

    def test_analysis_fallback_retains_specific_document_deliverables(self):
        agent = ObjectiveAgent.__new__(ObjectiveAgent)
        requirements = agent._get_fallback_requirements(
            "",
            "Module: Exception Triage\nDeliverable: BRD-17 maker-checker queue\n"
            "Acceptance Evidence: route unmatched and rejected instructions",
        )
        self.assertTrue(any(
            "BRD-17 maker-checker queue" in item
            for item in requirements["key_deliverables"]
        ))
        self.assertIn("Supporting document corpus", requirements["source_basis"])
        self.assertNotIn("generic AWS", requirements["project_overview"])

    def test_objective_quality_gate_rejects_unrelated_input_without_evidence(self):
        result = ObjectiveAgent.evaluate_input_quality({
            "input_assessment": {
                "is_sow_candidate": False, "confidence": 0.96,
                "reason": "The content is an unrelated personal conversation.",
            },
            "key_features": [], "functional_requirements": [], "key_deliverables": [],
            "workflow_steps": [], "use_cases": [], "architecture_components": [],
            "integration_details": [], "current_state": [],
        }, "What should I cook for dinner tonight?")
        self.assertFalse(result["accepted"])

    def test_objective_quality_gate_allows_a_short_clear_project_request(self):
        result = ObjectiveAgent.evaluate_input_quality({
            "input_assessment": {
                "is_sow_candidate": True, "confidence": 0.91,
                "reason": "A clear customer-support automation project is requested.",
            },
            "key_features": ["Customer-support chatbot"],
            "desired_outcomes": ["Reduce manual ticket handling"],
        }, "Build a customer-support chatbot with agent handover.")
        self.assertTrue(result["accepted"])

    def test_analysis_node_stops_after_objective_rejection(self):
        class RejectingObjectiveAgent:
            def __init__(self, _config):
                pass

            def analyze_objective(self, _objective, _supporting_context):
                return {"input_assessment": {"is_sow_candidate": False, "confidence": 0.99}}

            evaluate_input_quality = staticmethod(lambda *_args: {
                "accepted": False, "confidence": 0.99,
                "reason": "No project scope was supplied.", "evidence_categories": {},
            })

        with patch("app.core.nodes.ObjectiveAgent", RejectingObjectiveAgent):
            result = analyze_objective_node({
                "mode": "POC", "objective": "random unrelated content",
                "supporting_context": "", "metadata": {},
            })
        self.assertEqual(result["current_step"], "objective_rejected")
        self.assertTrue(result["errors"])

    def test_objective_json_parser_recovers_wrapped_json(self):
        parsed = ObjectiveAgent._parse_json(
            "Analysis follows:\n```json\n{\"key_features\":[\"Agent handover\"]}\n```\nDone"
        )
        self.assertEqual(parsed["key_features"], ["Agent handover"])

    def test_objective_json_parser_recovers_truncated_top_level_object(self):
        parsed = ObjectiveAgent._parse_json(
            '{"project_overview":"Customer support automation",'
            '"key_features":["Chatbot","Agent handover"]'
        )
        self.assertEqual(parsed["project_overview"], "Customer support automation")
        self.assertEqual(parsed["key_features"], ["Chatbot", "Agent handover"])

    def test_objective_json_parser_salvages_fields_before_partial_key(self):
        parsed = ObjectiveAgent._parse_json(
            '{"project_overview":"Customer support automation","key_feat'
        )
        self.assertEqual(parsed["project_overview"], "Customer support automation")

    def test_preview_endpoint_passes_uploaded_docx_evidence_into_agent_state(self):
        """Exercise the same multipart endpoint used by the browser UI."""
        from app.core import server

        captured = {}

        def capture_async(_preview_id, initial_state, _use_fast_mode):
            captured.update(initial_state)
            return None

        previous_testing = server.app.config.get("TESTING", False)
        previous_upload = server.app.config.get("UPLOAD_FOLDER")
        with tempfile.TemporaryDirectory() as upload_dir:
            server.app.config["TESTING"] = True
            server.app.config["UPLOAD_FOLDER"] = upload_dir
            try:
                with (
                    patch.object(server, "retrieve_rag_data", return_value=({}, False)) as retrieve_rag,
                    patch("app.preview.async_preview.process_preview_async", side_effect=capture_async),
                ):
                    response = server.app.test_client().post(
                        "/api/preview",
                        data={
                            "mode": "POC",
                            "company_name": "Axis Securities",
                            "author_name": "Arvind",
                            "project_name": "Settlement Operations",
                            "additional_details": "Keep voice capability as future scope.",
                            "business_unit": "Cloud",
                            "supporting_doc_count": "1",
                            "selected_sow_sections": json.dumps([
                                "project_overview", "scope_of_work"
                            ]),
                            "supporting_docs": (
                                self._build_brd_bytes(), "specific_brd.docx"
                            ),
                        },
                        content_type="multipart/form-data",
                    )
            finally:
                server.app.config["TESTING"] = previous_testing
                server.app.config["UPLOAD_FOLDER"] = previous_upload

        payload = response.get_json()
        self.assertEqual(response.status_code, 202, payload)
        self.assertEqual(payload["supporting_documents_received"], 1)
        self.assertEqual(payload["supporting_documents_extracted"], 1)
        self.assertGreater(payload["supporting_context_chars"], 100)
        self.assertEqual(captured["objective"], "Keep voice capability as future scope.")
        self.assertEqual(captured["additional_details"], "Keep voice capability as future scope.")
        self.assertEqual(captured["metadata"]["business_unit"], "Cloud")
        self.assertEqual(captured["metadata"]["owner_email"], "test-admin@shellkode.com")
        self.assertEqual(payload["metadata"]["business_unit"], "Cloud")
        retrieve_rag.assert_not_called()
        self.assertIn("specific_brd.docx", captured["supporting_context"])
        self.assertIn("BRD-17", captured["supporting_context"])
        self.assertIn("maker-checker assignment", captured["supporting_context"])


if __name__ == "__main__":
    unittest.main()
