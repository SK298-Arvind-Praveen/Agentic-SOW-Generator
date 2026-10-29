from pathlib import Path
from types import SimpleNamespace
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
import pytest

from app.core.section_formats import presentation_contract, timeline_format
from app.core.sow_section_preferences import SECTION_LABELS, section_category
from app.document.document_builder import SectionBuilder, DocumentBuilder
from app.core.sow_quality import validate_generated_sections


def test_contracts_are_source_only_and_format_specific():
    assert 'explicitly stated' in presentation_contract('success_criteria', {})
    assert 'leave the section body blank' in presentation_contract('success_criteria', {})
    assert 'milestone table' in presentation_contract('timelines_deliverables', {})
    assert 'Gantt grid' in presentation_contract('timelines_deliverables', {'timeline_format': 'gantt'})
    with pytest.raises(ValueError):
        timeline_format('unknown')


def test_timeline_name_and_legacy_content_are_compatible():
    assert SECTION_LABELS['timelines_deliverables'] == 'Project Timeline'
    assert section_category('Project Timeline') == section_category('Timelines and Deliverables')
    builder = DocumentBuilder.__new__(DocumentBuilder)
    assert builder._resolve_content('Project Timeline', {'timelines_and_deliverables': 'schedule'}) == 'schedule'


def test_no_source_backed_outcomes_can_leave_success_criteria_blank():
    missing, issues = validate_generated_sections({'success_criteria': ''}, 'POC', ['success_criteria'])
    assert not missing and not issues


@pytest.mark.parametrize('sentinel', ['[NO_SOURCE_BACKED_SUCCESS_CRITERIA]', '[NO_QUANTITATIVE_TARGETS]'])
def test_writer_removes_internal_no_targets_sentinel(monkeypatch, sentinel):
    from app.agents.poc_writer_agent import POCWriterAgent, TemplateSection
    agent = POCWriterAgent.__new__(POCWriterAgent)
    agent.config = SimpleNamespace()
    agent.template_type = 'POC'
    monkeypatch.setattr(agent, '_build_individual_prompt', lambda *a, **kw: 'prompt')
    monkeypatch.setattr(agent, '_section_token_budget', lambda *a: 1000)
    monkeypatch.setattr(agent, '_call_bedrock', lambda *a, **kw: sentinel)
    section = TemplateSection('Success Criteria', '', {}, 1)
    assert agent._generate_section(section, {}, {}, '', []) == ''


def test_success_contract_allows_factual_descriptions_without_inventing_targets():
    contract = presentation_contract('success_criteria', {})
    assert 'descriptive success criteria' in contract
    assert 'when no numerical targets exist' in contract
    assert 'Every criterion must be traceable to the BRD/user inputs' in contract
    assert 'Do not invent percentages' in contract
    assert 'Preserve conditions, exclusions and phase boundaries' in contract
    assert 'neither quantitative targets nor supported functional outcomes' in contract
    assert 'No descriptive outcomes' not in contract


@pytest.mark.parametrize('content', [
    '- Customers can retrieve order status through the web chatbot.',
    '- Reduce refresh time from 3 hours to 1 hour.\n- Transfer unresolved chats with their conversation history.',
])
def test_writer_retains_supported_functional_and_mixed_success_bullets(monkeypatch, content):
    from app.agents.poc_writer_agent import POCWriterAgent, TemplateSection
    agent = POCWriterAgent.__new__(POCWriterAgent)
    agent.config = SimpleNamespace()
    agent.template_type = 'POC'
    monkeypatch.setattr(agent, '_build_individual_prompt', lambda *a, **kw: 'prompt')
    monkeypatch.setattr(agent, '_section_token_budget', lambda *a: 1000)
    monkeypatch.setattr(agent, '_call_bedrock', lambda *a, **kw: content)
    section = TemplateSection('Success Criteria', '', {}, 1)
    output = agent._generate_section(section, {}, {}, content, [])
    assert output == content
    assert validate_generated_sections({'success_criteria': output}, 'POC', ['success_criteria']) == ([], [])


def test_gantt_stays_one_table_with_coloured_cells():
    doc = Document()
    doc.styles.add_style('SOW Table', WD_STYLE_TYPE.TABLE)
    builder = SectionBuilder(doc, SimpleNamespace())
    builder.build_table([
        '| S.No | Activity | M1 | M2 | M3 | M4 | M5 | M6 | M7 | M8 |',
        '|---|---|---|---|---|---|---|---|---|---|',
        '| 1 | Discovery | ■ | | | | | | | |',
        '| 2 | Support | | | | | | | | ◆ |',
    ])
    assert len(doc.tables) == 1 and len(doc.tables[0].columns) == 10
    table = doc.tables[0]
    fill = lambda row, col: table.cell(row, col)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill'))
    assert fill(0, 2) == 'D9D9D9'
    assert fill(1, 2) == 'CCFFCC'
    assert fill(2, 9) == '4A86E8'
    assert table.cell(1, 2).text == ''


@pytest.mark.parametrize('filename', ['poc_template.md', 'production_template.md', 'poc_to_prod_template.md'])
def test_every_template_has_formats_and_source_backed_success_contract(filename):
    text = (Path(__file__).parents[1] / 'templates' / filename).read_text(encoding='utf8')
    assert '## Project Timeline' in text
    assert 'GANTT FORMAT:' in text and 'TABLE FORMAT:' in text
    assert '## Success Criteria' in text and 'quantifiable success targets' in text
    assert 'descriptive success criteria' in text
    assert 'when no numerical targets exist' in text
    assert text.count('## Success Criteria') == 1
