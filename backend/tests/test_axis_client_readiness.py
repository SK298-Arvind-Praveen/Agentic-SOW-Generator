"""Regressions from the Axis v27 client-readiness review; no AWS calls."""
import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from app.agents.poc_writer_agent import POCWriterAgent
from app.agents.sow_reviewer_agent import SOWReviewerAgent, apply_bullet_emphasis, bullet_review_checks
from app.core.sow_editorial import SECTION_PURPOSES, EDITORIAL_CONTRACT
from app.document.document_builder import DocumentBuilder, SectionBuilder
from app.pricing.renderer import render_aws_pricing_section
from app.pricing.service import AwsPricingService


def test_architecture_figure_has_no_generated_commentary_and_preserves_section_bullets():
    builder = DocumentBuilder.__new__(DocumentBuilder)
    builder.doc = Document()
    data = io.BytesIO()
    Image.new('RGB', (80, 40), 'white').save(data, format='PNG')
    builder._add_architecture_diagram({'image_base64': base64.b64encode(data.getvalue()).decode(),
        'title': 'Generated title', 'description': 'Unwanted generated description.',
        'note': 'Hosting requires validation', 'edit_url': 'https://app.diagrams.net/',
        'alt_text': 'Accessible architecture'})
    builder.doc.add_paragraph('The integration workers process customer requests.', style='List Bullet')
    paragraphs = builder.doc.paragraphs
    assert paragraphs[0]._p.xpath('.//w:drawing')
    assert len(paragraphs) == 2
    assert paragraphs[1].style.name == 'List Bullet'
    assert paragraphs[1].text == 'The integration workers process customer requests.'
    assert 'Unwanted' not in '\n'.join(p.text for p in paragraphs)
    assert builder.doc.inline_shapes[0]._inline.docPr.get('descr') == 'Accessible architecture'


def test_team_effort_is_table_only_in_writer_and_legacy_export():
    table = '| Role | Count | Effort |\n|---|---|---|\n| Architect | | |'
    body = 'Introductory filler.\n\n' + table + '\n\nWe will assign everyone during Week 1.'
    writer = POCWriterAgent.__new__(POCWriterAgent)
    assert writer._clean_content(body, 'Shellkode Project Team Effort') == writer._clean_content(table, 'Shellkode Project Team Effort')
    builder = DocumentBuilder.__new__(DocumentBuilder)
    assert builder._resolve_section_content('14. Shellkode Project Team Effort',
        {'shellkode_implementation_cost': body}, {}) == table


def test_model_selected_emphasis_does_not_rewrite_words_or_tables():
    original = '| Item | Value |\n|---|---|\n| A | B |\n\n- Migrate historical records with linked attachments.\n- Obtain approved rules for configuration.\n'
    result = apply_bullet_emphasis(original, {'1': 'historical records', '2': 'approved rules'})
    assert result.replace('**', '') == original
    assert '**historical records**' in result and '**approved rules**' in result
    assert bullet_review_checks(result, 'assumptions')['missing_inline_emphasis'] == []
    document = Document()
    builder = DocumentBuilder.__new__(DocumentBuilder)
    builder.doc = document
    builder._configure_styles()
    SectionBuilder(document, SimpleNamespace()).parse_content(result)
    assert any(run.bold and 'historical records' in run.text for p in document.paragraphs for run in p.runs)


@pytest.mark.parametrize('highlights', [
    {'0': 'records'}, {'2': 'records'}, {'1': 'invented'}, {'1': 'cord'},
    {'1': 'Migrate records securely'}, {'1': '**records**'}, {'1': 'records\n'},
])
def test_unsafe_emphasis_is_rejected(highlights):
    with pytest.raises(ValueError):
        apply_bullet_emphasis('- Migrate records securely.', highlights)


@pytest.mark.parametrize('body,phrase', [
    ('- Use [client API](https://example.com) for orders.', 'client API'),
    ('- Connect to https://example.com/api for orders.', 'api'),
    ('- Store the `status` value.', 'status'),
])
def test_emphasis_cannot_damage_links_or_code(body, phrase):
    with pytest.raises(ValueError):
        apply_bullet_emphasis(body, {'1': phrase})


def test_reviewer_can_use_compact_emphasis_and_logs_unfixed_bullets(tmp_path):
    original = {'assumptions': '- Customer provides approved rules.\n- Historical exports retain linked identifiers.',
                'legal': 'Exact static text.'}
    def call(prompt):
        assert 'material project-specific delivery premises' in prompt
        assert 'missing_inline_emphasis' in prompt
        return json.dumps({'corrections': [{'section': 'assumptions', 'emphasis': {'1': 'approved rules'},
                                           'reason': 'Emphasise the delivery premise.'}]})
    reviewer = SOWReviewerAgent(SimpleNamespace(SOW_REVIEW_LOG_DIR=tmp_path), call)
    result = reviewer.review(original, {'assumptions': 'assumptions'}, {})
    assert result['legal'] == original['legal']
    assert '**approved rules**' in result['assumptions']
    audit = reviewer.last_review_audit
    assert audit['batches'][0]['changes'][0]['status'] == 'applied'
    assert audit['bullet_checks_after']['assumptions']['missing_inline_emphasis'] == [2]
    assert not audit['review_complete']


def test_emphasis_applies_after_content_edits(tmp_path):
    response = {'corrections': [
        {'section': 'benefits', 'after': '- Improve customer visibility.', 'reason': 'Remove repeated outcome.'},
        {'section': 'benefits', 'emphasis': {'1': 'customer visibility'}, 'reason': 'Highlight outcome.'}]}
    reviewer = SOWReviewerAgent(SimpleNamespace(SOW_REVIEW_LOG_DIR=tmp_path), lambda _: json.dumps(response))
    result = reviewer.review({'benefits': '- Improve visibility.\n- Improve visibility again.'},
                             {'benefits': 'business_benefits'}, {})
    assert result['benefits'] == '- Improve **customer visibility**.'


def test_pricing_does_not_treat_workflow_deadlines_as_workload():
    metrics = [('Regulatory TAT', '21 days'), ('Escalation days', '15,18,21'),
               ('Ticket reopen window', '3 days'), ('Inbox records per page', '50-500'),
               ('Monthly ticket volume', '10000'), ('History retention', '12 months'),
               ('Concurrent users', '50')]
    rendered = render_aws_pricing_section({'status': 'needs_input', 'volume_metrics':
                                         [{'metric': k, 'value': v} for k, v in metrics]})
    for label, _ in metrics[:4]:
        assert label not in rendered
    for label, value in metrics[4:]:
        assert f'| {label} | {value} |' in rendered


def test_workflow_rules_are_also_excluded_before_cost_estimation():
    plan = {'volume_metrics': [
        {'metric': 'regulatory TAT', 'value': '21 days', 'evidence_quote': 'regulatory TAT 21 days'},
        {'metric': 'Monthly tickets', 'value': '10000', 'evidence_quote': 'Monthly tickets 10000'}]}
    rows = AwsPricingService._validated_volume_metrics(plan, 'regulatory TAT 21 days. Monthly tickets 10000.')
    assert [row['metric'] for row in rows] == ['Monthly tickets']


def test_reviewer_sees_bounded_original_evidence_not_only_model_interpretation(tmp_path):
    source = 'Email compose supports grammar checking.\n\n' + 'Other details.\n\n' * 1200
    def call(prompt):
        assert 'Email compose supports grammar checking.' in prompt
        assert '_review_source_text' not in prompt
        excerpt = prompt.split('ORIGINAL EVIDENCE EXCERPTS', 1)[1].split('Use only these exact section keys', 1)[0]
        assert len(excerpt) < 4200
        assert 'not proof of customer agreement' in prompt
        return '{"corrections":[]}'
    reviewer = SOWReviewerAgent(SimpleNamespace(SOW_REVIEW_LOG_DIR=tmp_path), call)
    reviewer.review({'excluded': '- Writing assistance is excluded.'}, {'excluded': 'out_of_scope'},
                    {'_review_source_text': source})
    assert reviewer.last_review_audit['batches'][0]['status'] == 'completed'


def test_all_default_templates_use_material_delivery_premises():
    assert 'approved business/regulatory rules' in SECTION_PURPOSES['assumptions']
    assert 'Connect CRM is not automatically Amazon Connect' in EDITORIAL_CONTRACT
    for name in ('poc', 'production', 'poc_to_prod'):
        text = (Path(__file__).resolve().parents[1] / 'templates' / f'{name}_template.md').read_text(encoding='utf-8')
        assert 'State only material project-specific delivery premises' in text
        assert 'In every generated narrative bullet' in text


def test_identity_rules_reach_analysis_and_diagram_planning():
    from app.agents.objective_agent import ObjectiveAgent
    from app.diagram.planning import ARCHITECTURE_PLANNING_RULES
    analyst = ObjectiveAgent.__new__(ObjectiveAgent)
    prompt = analyst._build_analysis_prompt('Migrate I-Leverage to Connect CRM.')
    assert 'Connect CRM is not automatically' in prompt
    assert 'NPA to' in prompt and 'NPS' in prompt
    assert 'does not mean Amazon Connect' in ARCHITECTURE_PLANNING_RULES
    assert 'disconnected decorative boxes' in ARCHITECTURE_PLANNING_RULES
