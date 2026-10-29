from pathlib import Path
from types import SimpleNamespace

from docx import Document

from app.agents.poc_writer_agent import POCWriterAgent, TemplateSection
from app.core.sow_section_preferences import section_category
from app.document.document_builder import DocumentBuilder, SectionBuilder


def test_current_state_is_never_implicitly_selected_with_objective():
    assert section_category('5. Current State') == 'current_state'
    assert section_category('Objective') == 'project_overview'
    writer = POCWriterAgent.__new__(POCWriterAgent)
    sections = [TemplateSection('Objective', '', {}, 1), TemplateSection('Current State', '', {}, 2)]
    selected = writer._order_sections_by_preference(sections, ['project_overview'])
    assert [s.name for s in selected] == ['Objective']
    assert writer._order_sections_by_preference(sections, ['current_state']) == [sections[1]]
    for filename in ('poc_template.md', 'production_template.md'):
        template = (Path(__file__).parents[1] / 'templates' / filename).read_text(encoding='utf8')
        assert '## Current State' not in template


def test_selected_empty_success_criteria_has_real_heading_and_bookmark():
    builder = DocumentBuilder.__new__(DocumentBuilder)
    builder.doc = Document()
    builder.config = SimpleNamespace()
    builder.section_builder = SectionBuilder(builder.doc, builder.config)
    builder._bookmark_ids = 0
    builder.subheading_anchor_maps = {}
    assert builder._preserve_blank_section('17. Success Criteria', {'success_criteria': ''})
    assert not builder._preserve_blank_section('17. Success Criteria', {})
    builder._build_section('17. Success Criteria', {'success_criteria': ''}, {}, 16)
    assert builder.doc.paragraphs[0].text == '17. Success Criteria'
    assert 'bookmarkStart' in builder.doc.paragraphs[0]._p.xml
    assert 'NO_QUANTITATIVE' not in builder.doc._element.xml


def test_blank_source_only_success_criteria_is_not_sent_for_review(monkeypatch):
    from app.agents.sow_reviewer_agent import SOWReviewerAgent
    def unexpected(*args, **kwargs):
        raise AssertionError('No review call should be made for intentional empty criteria')
    monkeypatch.setattr(SOWReviewerAgent, 'review', unexpected)
    writer = POCWriterAgent.__new__(POCWriterAgent)
    writer.config = SimpleNamespace()
    writer._review_document_relevance({'success_criteria': ''},
                                     [TemplateSection('Success Criteria', '', {}, 1)], {})
