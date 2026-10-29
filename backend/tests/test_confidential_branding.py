from pathlib import Path
from types import SimpleNamespace
from docx import Document
from docx.oxml.ns import qn
from app.document.document_builder import DocumentBuilder, SectionBuilder
from app.core.sow_editorial import EDITORIAL_CONTRACT


def test_confidential_footer_and_table_header_colour():
    builder = DocumentBuilder.__new__(DocumentBuilder)
    builder.doc = Document()
    builder.config = SimpleNamespace(ASSETS_DIR=Path(__file__).resolve().parents[1] / 'assets')
    builder._configure_styles()
    builder._add_header_footer(builder.doc.sections[0], {}, 'POC')
    footer = builder.doc.sections[0].footer
    paragraph = next(p for p in footer.paragraphs if 'limited purpose' in p.text)
    assert paragraph.text.startswith('Shellkode Confidential - for limited purpose only')
    assert paragraph.paragraph_format.alignment == 2
    assert paragraph.runs[0].font.size.pt == 8
    assert paragraph._p.xpath('./w:pPr/w:pBdr/w:top')
    assert all(not p.text for p in builder.doc.sections[0].header.paragraphs)
    assert not builder.doc.sections[0].header.tables
    assert 'Copyright' not in '\n'.join(p.text for p in footer.paragraphs)
    assert str(paragraph.runs[0].font.color.rgb) == '7F00FF'
    table = SectionBuilder(builder.doc, builder.config)._build_table_rows([
        ['Item', 'Monthly (USD)', 'Estimate'], ['Infrastructure', '100', 'Planning estimate']])
    for cell in table.rows[0].cells:
        assert cell._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')) == '7F00FF'
        assert str(cell.paragraphs[0].runs[0].font.color.rgb) == 'FFFFFF'


def test_short_bullet_contract_and_company_fallback():
    from app.agents.poc_writer_agent import POCWriterAgent
    assert '8-16 words' in EDITORIAL_CONTRACT
    assert 'two brief factual paragraphs' in EDITORIAL_CONTRACT
    text = POCWriterAgent._about_company_fallback({'company_name': 'Example'})
    assert not text.startswith('- ')
