import base64
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.diagram.service import (
    ASSET_KEY,
    DiagramService,
    _geometry,
    drawio_xml,
    edit_url,
    fallback_spec,
    update_asset,
    validate_drawio_xml,
    validate_spec,
)
from app.document.document_builder import DocumentBuilder, SectionBuilder


class _Body:
    def __init__(self, value):
        self.value = value

    def read(self):
        return json.dumps(self.value).encode("utf-8")


class _Bedrock:
    def __init__(self, text):
        self.text = text

    def invoke_model(self, **_kwargs):
        return {"body": _Body({"content": [{"text": self.text}], "usage": {}})}


class _Config:
    MODEL_ID = "test-model"

    def __init__(self, output_dir=None):
        backend = Path(__file__).resolve().parents[1]
        self.ASSETS_DIR = backend / "assets"
        self.COVER_PAGE_IMAGE = self.ASSETS_DIR / "coverpage.png"
        self.OUTPUT_DIR = Path(output_dir or tempfile.gettempdir())


class ArchitectureDiagramTests(unittest.TestCase):
    def test_validated_spec_compiles_to_editable_xml_and_png(self):
        response = {
            "title": "CRM Integration Architecture",
            "nodes": [
                {"id": "users", "label": "Advisers", "kind": "actor", "layer": 0},
                {"id": "crm", "label": "CRM Portal", "kind": "application", "layer": 1},
                {"id": "lambda", "label": "AWS Lambda", "kind": "service", "layer": 2},
            ],
            "edges": [
                {"source": "users", "target": "crm", "label": "HTTPS"},
                {"source": "crm", "target": "lambda", "label": "API"},
            ],
        }
        service = DiagramService(_Config(), bedrock=_Bedrock(json.dumps(response)))
        asset = service.generate_asset(
            {"aws_services": ["AWS Lambda"]},
            {"company_name": "Axis Securities", "project_title": "CRM Integration"},
            "The portal invokes AWS Lambda through an API.",
        )
        self.assertFalse(asset["used_fallback"])
        self.assertEqual(asset["type"], "drawio_architecture")
        self.assertIn("<mxfile", asset["drawio_xml"])
        self.assertIn("AWS Lambda", asset["drawio_xml"])
        self.assertTrue(base64.b64decode(asset["image_base64"]).startswith(b"\x89PNG"))
        self.assertTrue(asset["edit_url"].startswith("https://app.diagrams.net/"))

    def test_edit_url_uses_external_query_instead_of_word_bookmark_fragment(self):
        url = edit_url("<mxGraphModel><root /></mxGraphModel>")
        self.assertIn("&create=", url)
        self.assertNotIn("#create=", url)

    def test_evidence_gate_plans_multiple_distinct_diagrams_without_exceeding_cap(self):
        diagrams = []
        for diagram_type, title in (
            ("architecture_overview", "CRM Architecture Overview"),
            ("integration_context", "External Integration Context"),
            ("data_flow", "Customer Request Data Flow"),
            ("deployment_topology", "Unrequested Extra Diagram"),
        ):
            diagrams.append({
                "diagram_type": diagram_type,
                "title": title,
                "placement_heading": "End-to-End Data Flow" if diagram_type == "data_flow" else "Integrations",
                "nodes": [
                    {"id": "source", "label": "Source", "kind": "external", "layer": 0},
                    {"id": "solution", "label": "CRM", "kind": "application", "layer": 1},
                    {"id": "target", "label": "Amazon Bedrock", "kind": "service", "layer": 2},
                ],
                "edges": [
                    {"source": "source", "target": "solution", "label": "Request"},
                    {"source": "solution", "target": "target", "label": "Inference"},
                ],
            })
        service = DiagramService(_Config(), bedrock=_Bedrock(json.dumps({"diagrams": diagrams})))
        assets = service.generate_assets(
            {
                "aws_services": ["Amazon Bedrock"],
                "integrations": ["Cisco", "Netcore", "Active Directory", "Core API"],
            },
            {"company_name": "Axis Securities", "project_title": "CRM Integration"},
            "The request workflow includes ingestion, data flow, and a response to the agent.",
        )
        self.assertEqual(len(assets), 3)
        self.assertEqual(
            [asset["diagram_type"] for asset in assets],
            ["architecture_overview", "integration_context", "data_flow"],
        )
        self.assertTrue(all(asset["caption"] == asset["title"] for asset in assets))
        self.assertTrue(all(asset["description"] for asset in assets))

    def test_flow_semantics_render_with_horizontal_ranks_and_standard_shapes(self):
        raw = {
            "title": "Approval Flow",
            "description": "This diagram depicts request review and approval outcomes.",
            "nodes": [
                {"id": "start", "label": "Request received", "kind": "terminator", "layer": 0},
                {"id": "capture", "label": "Capture request", "kind": "input_output", "layer": 0},
                {"id": "review", "label": "Review request", "kind": "process", "layer": 0},
                {"id": "approved", "label": "Approved?", "kind": "decision", "layer": 0},
                {"id": "end", "label": "Complete", "kind": "terminator", "layer": 0},
            ],
            "edges": [
                {"source": "start", "target": "capture"},
                {"source": "capture", "target": "review"},
                {"source": "review", "target": "approved"},
                {"source": "approved", "target": "end", "label": "Yes"},
            ],
        }
        service = DiagramService(_Config(), bedrock=_Bedrock("{}"))
        enriched = service._enrich_diagram(raw, "data_flow", {})
        spec = validate_spec(enriched, "Approval Flow")
        positions, width, height = _geometry(spec)
        xml = drawio_xml(spec)
        self.assertGreater(len({position[0] for position in positions.values()}), 1)
        self.assertGreater(width, height)
        self.assertIn("shape=rhombus", xml)
        self.assertIn("shape=parallelogram", xml)
        self.assertIn("arcSize=50", xml)

    def test_layout_caps_aspect_ratio_and_number_of_columns_for_readability(self):
        raw = {
            "title": "A deliberately long solution architecture title that must remain readable",
            "nodes": [
                {"id": f"node_{index}", "label": f"Component {index}", "kind": "service", "layer": index}
                for index in range(6)
            ],
            "edges": [
                {"source": f"node_{index}", "target": f"node_{index + 1}", "label": "Authenticated request"}
                for index in range(5)
            ],
        }
        spec = validate_spec(raw, "Architecture")
        positions, width, height = _geometry(spec)
        self.assertLessEqual(width / height, 2.0)
        self.assertLessEqual(len({position[0] for position in positions.values()}), 4)

    def test_long_linear_flow_wraps_as_a_compact_serpentine(self):
        raw = {
            "title": "Case Flow",
            "nodes": [
                {"id": f"step_{index}", "label": f"Step {index}", "kind": "process", "layer": index}
                for index in range(9)
            ],
            "edges": [
                {"source": f"step_{index}", "target": f"step_{index + 1}"}
                for index in range(8)
            ],
        }
        spec = validate_spec(raw, "Case Flow")
        positions, width, height = _geometry(spec)
        self.assertGreater(width, height)
        self.assertEqual(positions["step_3"][0], positions["step_4"][0])
        self.assertGreater(positions["step_4"][1], positions["step_3"][1])
        self.assertGreater(positions["step_4"][0], positions["step_5"][0])

    def test_branch_heavy_flow_uses_bounded_landscape_grid(self):
        nodes = [{"id": "start", "label": "Start", "kind": "terminator", "layer": 0}]
        nodes.extend(
            {"id": f"task_{index}", "label": f"Task {index}", "kind": "process", "layer": 1}
            for index in range(8)
        )
        edges = [{"source": "start", "target": f"task_{index}"} for index in range(8)]
        spec = validate_spec({"title": "Branched Flow", "nodes": nodes, "edges": edges}, "Branched Flow")
        positions, width, height = _geometry(spec)
        self.assertGreater(width, height)
        self.assertLessEqual(len({position[1] for position in positions.values()}), 3)

    def test_integration_peer_chain_is_rewritten_as_hub_and_spoke(self):
        nodes = [
            {"id": "platform", "label": "Customer Support Platform", "kind": "application", "layer": 1},
            {"id": "hybris", "label": "SAP Hybris", "kind": "external", "layer": 0},
            {"id": "genesys", "label": "Genesys", "kind": "external", "layer": 2},
            {"id": "redshift", "label": "Redshift", "kind": "external", "layer": 3},
            {"id": "whatsapp", "label": "WhatsApp", "kind": "channel", "layer": 4},
        ]
        edges = [
            {"source": "hybris", "target": "genesys"},
            {"source": "genesys", "target": "platform"},
            {"source": "platform", "target": "redshift"},
            {"source": "redshift", "target": "whatsapp"},
        ]
        rewritten_nodes, rewritten_edges = DiagramService._normalise_integration_topology(nodes, edges)
        self.assertEqual(len(rewritten_edges), 4)
        self.assertTrue(all("platform" in {edge["source"], edge["target"]} for edge in rewritten_edges))
        spec = validate_spec({"title": "Integrations", "nodes": rewritten_nodes, "edges": rewritten_edges}, "Integrations")
        positions, width, height = _geometry(spec)
        self.assertGreater(width, height)

    def test_invalid_agent_output_uses_deterministic_fallback(self):
        service = DiagramService(_Config(), bedrock=_Bedrock("not json"))
        asset = service.generate_asset(
            {"aws_services": ["Amazon Bedrock", "Amazon DynamoDB"]},
            {"company_name": "Example", "project_title": "Support Assistant"},
            "",
        )
        self.assertTrue(asset["used_fallback"])
        self.assertTrue(base64.b64decode(asset["image_base64"]).startswith(b"\x89PNG"))

    def test_network_role_and_following_flow_both_render(self):
        response = {"diagrams": [
            {
                "diagram_type": "architecture_overview",
                "title": "Migration Architecture",
                "nodes": [
                    {"id": "source", "label": "On-Premises", "kind": "external", "layer": 0,
                     "placement_role": "external"},
                    {"id": "vpn", "label": "VPN", "kind": "service", "layer": 1,
                     "placement_role": "entry"},
                    {"id": "firewall", "label": "Firewall", "kind": "service", "layer": 1,
                     "placement_role": "network"},
                    {"id": "workload", "label": "Workload", "kind": "service", "layer": 2,
                     "placement_role": "compute"},
                ],
                "edges": [
                    {"source": "source", "target": "vpn"},
                    {"source": "vpn", "target": "firewall"},
                    {"source": "firewall", "target": "workload"},
                ],
            },
            {
                "diagram_type": "data_flow",
                "title": "Migration Flow",
                "nodes": [
                    {"id": "start", "label": "Start", "kind": "terminator", "layer": 0},
                    {"id": "copy", "label": "Transfer", "kind": "process", "layer": 1},
                    {"id": "end", "label": "Complete", "kind": "terminator", "layer": 2},
                ],
                "edges": [
                    {"source": "start", "target": "copy"},
                    {"source": "copy", "target": "end"},
                ],
            },
        ]}
        service = DiagramService(_Config(), bedrock=_Bedrock(json.dumps(response)))
        assets = service.generate_assets(
            {"aws_services": ["AWS Site-to-Site VPN"], "workflow_steps": ["Transfer data"]},
            {"company_name": "Example", "project_title": "Cloud Migration"},
            "The migration workflow transfers workloads through a firewall and validates the data flow.",
        )
        self.assertEqual([asset["diagram_type"] for asset in assets], ["architecture_overview", "data_flow"])
        self.assertTrue(all(base64.b64decode(asset["image_base64"]).startswith(b"\x89PNG") for asset in assets))

    def test_drawio_edit_round_trip_validates_png_and_xml(self):
        spec = fallback_spec(
            {"aws_services": ["Amazon S3"]},
            {"company_name": "Example", "project_title": "Document Workflow"},
        )
        xml = drawio_xml(spec)
        service = DiagramService(_Config(), bedrock=_Bedrock("{}"))
        image_base64 = service.generate_asset(
            {"aws_services": ["Amazon S3"]},
            {"company_name": "Example", "project_title": "Document Workflow"},
            "",
        )["image_base64"]
        png = "data:image/png;base64," + image_base64
        updated = update_asset(xml, png, {"title": "Existing"})
        self.assertEqual(validate_drawio_xml(updated["drawio_xml"]), xml)
        self.assertTrue(updated["edit_url"].startswith("https://app.diagrams.net/"))
        with self.assertRaises(ValueError):
            update_asset("<html />", png)
        with self.assertRaises(ValueError):
            update_asset("<mxfile><broken></mxfile>", png)

    def test_document_section_embeds_image_and_external_edit_link(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config = _Config(temp_dir)
            service = DiagramService(config, bedrock=_Bedrock("invalid"))
            asset = service.generate_asset(
                {"aws_services": ["Amazon Bedrock"]},
                {"company_name": "Example", "project_title": "AI Assistant"},
                "",
            )
            builder = DocumentBuilder(config)
            builder.doc = Document()
            builder._configure_styles()
            builder._add_numbering()
            builder.section_builder = SectionBuilder(builder.doc, config)
            builder._build_section(
                "1. Solution Architecture (AWS)",
                {
                    "architecture_diagram": "The proposed logical design is shown below.",
                    "architecture_diagram_asset": asset,
                },
                {"company_name": "Example", "author_org": "ShellKode"},
                0,
            )
            self.assertEqual(len(builder.doc.inline_shapes), 1)
            self.assertFalse(any(
                paragraph.text.startswith("Figure 1:")
                for paragraph in builder.doc.paragraphs
            ))
            relationship_targets = [rel.target_ref for rel in builder.doc.part.rels.values()]
            self.assertIn(asset["edit_url"], relationship_targets)
            self.assertTrue(any(
                "This diagram depicts" in paragraph.text
                for paragraph in builder.doc.paragraphs
            ))

    def test_non_overview_diagrams_are_assigned_to_matching_sow_sections(self):
        builder = DocumentBuilder(_Config())
        builder.toc_entries = [
            "1. Solution Architecture",
            "2. Current Workflow and Pain Points",
            "3. Integrations and Dependencies",
        ]
        assets = [
            {"diagram_type": "architecture_overview", "title": "Solution Architecture"},
            {"diagram_type": "data_flow", "title": "Request Workflow", "placement_heading": "Current Workflow"},
            {"diagram_type": "integration_context", "title": "System Integrations", "placement_heading": "Integrations"},
        ]
        sections = {
            "solution_architecture": "AWS design.",
            "current_workflow_and_pain_points": "The current request workflow.",
            "integrations_and_dependencies": "External system integrations.",
        }
        assigned = builder._assign_diagram_sections(assets, sections, {})
        self.assertEqual(assigned[id(assets[0])], builder.toc_entries[0])
        self.assertEqual(assigned[id(assets[1])], builder.toc_entries[1])
        self.assertEqual(assigned[id(assets[2])], builder.toc_entries[2])

    def test_preview_endpoint_persists_an_edited_diagram(self):
        from app.core import server

        preview_id = "PREVIEW_TEST_DIAGRAM"
        service = DiagramService(_Config(), bedrock=_Bedrock("invalid"))
        assets = service.generate_assets(
            {"aws_services": ["Amazon S3"]},
            {"company_name": "Example", "project_title": "Document Workflow"},
            "",
        )
        with server.preview_lock:
            server.preview_storage[preview_id] = {
                "status": "ready",
                "metadata": {"business_unit": "GenAI"},
                "content": {ASSET_KEY: assets},
                "edit_count": 0,
            }
        try:
            server.app.config["TESTING"] = True
            response = server.app.test_client().put(
                f"/api/preview/{preview_id}/architecture-diagram",
                json={
                    "drawio_xml": assets[0]["drawio_xml"],
                    "image_data": "data:image/png;base64," + assets[0]["image_base64"],
                    "diagram_index": 0,
                },
            )
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.get_json()["success"])
            with server.preview_lock:
                self.assertEqual(server.preview_storage[preview_id]["edit_count"], 1)
        finally:
            server.app.config["TESTING"] = False
            with server.preview_lock:
                server.preview_storage.pop(preview_id, None)


if __name__ == "__main__":
    unittest.main()
