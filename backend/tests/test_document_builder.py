import re
import os
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from app.document.document_builder import (
    CHROME_INSET_IN,
    HORIZONTAL_MARGIN_IN,
    PAGE_WIDTH_IN,
    DocumentBuilder,
    SectionBuilder,
    _add_field,
    _bookmark_name,
    _find_libreoffice_binary,
    _flatten_pageref_fields,
    _pageref_results_from_xml,
)


class _Config:
    def __init__(self, output_dir):
        backend = Path(__file__).resolve().parents[1]
        self.OUTPUT_DIR = Path(output_dir)
        self.ASSETS_DIR = backend / "assets"
        self.TEMPLATES_DIR = backend / "templates"
        self.COVER_PAGE_TEMPLATE = self.TEMPLATES_DIR / "sow_coverpage_template.docx"
        self.COVER_PAGE_IMAGE = backend / "assets" / "coverpage.png"
        self.PRECOMPUTE_DOCUMENT_FIELDS = False


class DocumentBuilderTests(unittest.TestCase):
    def test_plain_external_url_is_rendered_as_clickable_hyperlink(self):
        document = Document()
        SectionBuilder(document, SimpleNamespace()).parse_content(
            "**AWS Pricing Calculator Link:** https://calculator.aws/#/estimate?id=abc"
        )
        hyperlink_relationships = [
            relationship.target_ref
            for relationship in document.part.rels.values()
            if relationship.reltype.endswith("/hyperlink")
        ]
        self.assertIn("https://calculator.aws/#/estimate?id=abc", hyperlink_relationships)

    def test_open_clarifications_table_reserves_readable_area_column(self):
        widths = SectionBuilder._column_widths([
            ["Module/Area", "Open Item", "Status / Note"],
            ["Performance Requirements", "Confirm concurrent users", "Required for sizing"],
        ])
        self.assertEqual(widths, [2450, 5000, 2832])

    def test_bookmark_names_stay_within_word_limit_and_remain_unique(self):
        title = "10.2 Effort Basis and Key Assumptions Affecting Delivery"
        first = _bookmark_name(title, 1034)
        second = _bookmark_name(title, 1035)
        self.assertLessEqual(len(first), 40)
        self.assertLessEqual(len(second), 40)
        self.assertNotEqual(first, second)
        self.assertTrue(first.startswith("SOW_1034_"))

    def test_calculated_pageref_can_be_materialised_for_google_docs(self):
        document = Document()
        paragraph = document.add_paragraph("Scope of Work")
        paragraph.add_run("\t")
        _add_field(paragraph, "PAGEREF SOW_1_Scope \\h", "7", size=11)
        flattened = _flatten_pageref_fields(document._element.xml.encode("utf-8"))
        xml = flattened.decode("utf-8")
        self.assertNotIn("PAGEREF", xml)
        self.assertIn(">7</w:t>", xml)
        self.assertIn("Scope of Work", xml)

    def test_finds_bundled_libreoffice_when_it_is_not_on_path(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_home = Path(temp_dir)
            bundled = (
                fake_home
                / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
            )
            bundled.parent.mkdir(parents=True)
            bundled.write_text("#!/bin/sh\n", encoding="utf-8")
            bundled.chmod(0o755)
            with patch("app.document.document_builder.shutil.which", return_value=None), patch(
                "app.document.document_builder.Path.home", return_value=fake_home
            ), patch.dict(
                os.environ,
                {
                    "ProgramFiles": str(fake_home / "Program Files"),
                    "ProgramFiles(x86)": str(fake_home / "Program Files (x86)"),
                },
            ):
                self.assertEqual(
                    _find_libreoffice_binary(_Config(temp_dir)), str(bundled)
                )

    def test_builder_preserves_order_semantics_and_geometry(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            builder = DocumentBuilder(_Config(temp_dir))
            sections = {
                "toc_structure": (
                    "1. Project Overview\n"
                    "2. Detailed Scope of Work\n"
                    "3. About Example Customer"
                ),
                "project_overview": (
                    "### Business Need\nA concise source-grounded overview.\n\n"
                    "- Confirm the selected workflow.\n"
                    "  - Record review evidence.\n"
                    "- Preserve the agreed boundary.\n\n"
                    "| Field | Value |\n|---|---|\n| Mode | POC |\n\n"
                    "Text after the first table must remain visible.\n\n"
                    "| ID | Evidence | Status |\n|---|---|---|\n| SC-01 | Test report | Proposed |"
                ),
                "scope_of_work": (
                    "### Module A\n#### Workflow\n1. Receive input.\n2. Validate input.\n\n"
                    "### Module B\n#### Workflow\n1. Process input.\n2. Record evidence.\n\n"
                    "### Effort Basis and Key Assumptions Affecting Delivery\n"
                    "A deliberately long TOC heading exercises the Word bookmark limit."
                ),
            }
            metadata = {
                "company_name": "Example Customer",
                "project_title": "Agentic CRM Platform",
                "author_name": "Test Author",
                "author_org": "ShellKode",
                "document_date": "18 August 2026",
                "version": "1.0",
            }
            output = builder.build_document(sections, metadata, mode="POC")
            self.assertIsNotNone(output)
            self.assertTrue(Path(output).exists())

            document = Document(output)
            self.assertEqual(document.styles["Normal"].font.size.pt, 11.0)
            self.assertEqual(document.styles["Heading 1"].font.size.pt, 20.0)
            self.assertEqual(document.styles["Heading 2"].font.size.pt, 15.0)
            self.assertEqual(document.styles["Heading 3"].font.size.pt, 15.0)
            self.assertEqual(document.styles["Heading 4"].font.size.pt, 15.0)
            self.assertEqual(document.styles["Normal"].paragraph_format.alignment, 3)
            self.assertEqual(document.styles["Normal"].paragraph_format.space_after.pt, 8.0)
            self.assertAlmostEqual(
                document.styles["Normal"].paragraph_format.line_spacing,
                1.25,
                places=2,
            )
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Table of Contents", text)
            toc_title = next(
                paragraph for paragraph in document.paragraphs
                if paragraph.text == "Table of Contents"
            )
            self.assertNotEqual(toc_title.style.name, "Heading 1")
            self.assertIn("Text after the first table must remain visible.", text)
            paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
            toc_index = paragraph_texts.index("Table of Contents")
            first_body_index = paragraph_texts.index("1. Project Overview")
            toc_text = "\n".join(
                paragraph.text for paragraph in document.paragraphs[toc_index:first_body_index]
            )
            self.assertIn("Business Need", toc_text)
            self.assertIn("Module A", toc_text)
            self.assertIn("Module B", toc_text)
            self.assertNotIn("About Example Customer", toc_text)
            toc_paragraphs = document.paragraphs[toc_index + 1:first_body_index]
            subtopic_rows = [
                paragraph for paragraph in toc_paragraphs
                if any(label in paragraph.text for label in ("Business Need", "Module A", "Module B"))
            ]
            self.assertTrue(subtopic_rows)
            self.assertTrue(all("\t" in paragraph.text for paragraph in subtopic_rows))
            self.assertGreaterEqual(len(document.tables), 2)
            self.assertTrue(all(table.style.name == "SOW Table" for table in document.tables))
            for table in document.tables:
                spacer = table._tbl.getnext()
                self.assertEqual(spacer.tag, qn("w:p"))
                spacing = spacer.find(qn("w:pPr")).find(qn("w:spacing"))
                self.assertEqual(spacing.get(qn("w:after")), "100")
            self.assertTrue(all(
                paragraph.alignment == 0
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            ))
            self.assertFalse(any(
                re.match(r"^(Table|Figure) \d+:", paragraph.text)
                for paragraph in document.paragraphs
            ))
            self.assertTrue(all(
                run.font.size and run.font.size.pt == 11.0
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
                for run in paragraph.runs
            ))
            self.assertTrue(any(p.style.name == "Heading 2" and p.text == "1.1 Business Need" for p in document.paragraphs))
            self.assertTrue(any(p.style.name == "Heading 2" and p.text == "2.1 Module A" for p in document.paragraphs))
            self.assertTrue(any(p.style.name == "Heading 3" and p.text == "2.1.1 Workflow" for p in document.paragraphs))
            subsection_headings = [
                paragraph for paragraph in document.paragraphs
                if paragraph.style.name in {"Heading 2", "Heading 3", "Heading 4"}
            ]
            self.assertTrue(subsection_headings)
            self.assertTrue(all(
                run.font.size and run.font.size.pt == 15.0
                for paragraph in subsection_headings
                for run in paragraph.runs
                if run.text
            ))

            numbered = [
                paragraph for paragraph in document.paragraphs
                if paragraph.text in {"Receive input.", "Validate input.", "Process input.", "Record evidence."}
            ]
            num_ids = [
                paragraph._p.pPr.numPr.numId.val
                for paragraph in numbered
            ]
            self.assertEqual(num_ids, [91, 91, 91, 91])

            bullets = [
                paragraph for paragraph in document.paragraphs
                if paragraph.text in {
                    "Confirm the selected workflow.",
                    "Record review evidence.",
                    "Preserve the agreed boundary.",
                }
            ]
            self.assertEqual(
                [paragraph._p.pPr.numPr.numId.val for paragraph in bullets],
                [91, 91, 91],
            )
            self.assertEqual(
                [paragraph._p.pPr.numPr.ilvl.val for paragraph in bullets],
                [0, 1, 0],
            )
            self.assertTrue(all(paragraph.alignment == 0 for paragraph in bullets))

            section = document.sections[-1]
            self.assertEqual(round(section.page_width.inches, 2), 8.5)
            self.assertEqual(round(section.page_height.inches, 2), 11.0)
            cover_section = document.sections[0]
            self.assertEqual(round(cover_section.page_width.inches, 2), 8.27)
            self.assertEqual(round(cover_section.page_height.inches, 2), 11.69)

            # Dynamic cover labels live in modern page-anchored DrawingML text
            # boxes. They remain editable w:t values without relying on flow
            # paragraphs or renderer-dependent frame positioning.
            cover_xml = document._element.xml
            self.assertIn("Example Customer", cover_xml)
            self.assertIn("Agentic CRM Platform", cover_xml)
            self.assertIn("Test Author", cover_xml)
            self.assertNotIn("Client Name", paragraph_texts)
            self.assertNotIn("<Author Name>", paragraph_texts)
            self.assertEqual(cover_xml.count("<wps:wsp>"), 6)
            self.assertNotIn("w:framePr", cover_xml)
            self.assertIn('name="Cover Company"', cover_xml)
            self.assertIn('<w:sz w:val="64"', cover_xml)
            self.assertIn('<w:sz w:val="30"', cover_xml)
            self.assertIn('<w:sz w:val="28"', cover_xml)
            self.assertIn('<w:sz w:val="24"', cover_xml)
            cover_boundary = next(
                paragraph for paragraph in document.paragraphs
                if paragraph._p.pPr is not None and paragraph._p.pPr.sectPr is not None
            )
            self.assertTrue(cover_boundary._p.xpath(".//w:drawing"))
            self.assertEqual(len(document.inline_shapes), 0)

            with zipfile.ZipFile(output) as package:
                document_xml_bytes = package.read("word/document.xml")
                xml = document_xml_bytes.decode("utf-8")
                numbering = package.read("word/numbering.xml").decode("utf-8")
            self.assertIn("w:tblHeader", xml)
            self.assertIn('<w:tblW', xml)
            self.assertIn('w:w="10282"', xml)
            self.assertIn('w:numId="91"', numbering)
            self.assertNotIn('w:numId="92"', numbering)
            self.assertIn("1.1 Business Need", xml)
            self.assertEqual(xml.count('TOC \\o "1-4" \\h \\z \\u'), 1)
            pageref_targets = re.findall(r"PAGEREF\s+([^\s<]+)\s+\\h", xml)
            bookmark_names = set(re.findall(r'<w:bookmarkStart[^>]+w:name="([^"]+)"', xml))
            self.assertTrue(pageref_targets)
            toc_rows = [
                paragraph
                for paragraph in document.paragraphs[toc_index + 1:first_body_index]
                if paragraph.text.strip()
            ]
            self.assertEqual(len(pageref_targets), len(toc_rows))
            self.assertTrue(set(pageref_targets).issubset(bookmark_names))
            cached_page_numbers = _pageref_results_from_xml(document_xml_bytes)
            self.assertTrue(cached_page_numbers)
            self.assertTrue(all(not value or value.isdigit() for value in cached_page_numbers.values()))

    def test_singular_timeline_title_resolves_shared_timeline_content_key(self):
        builder = DocumentBuilder.__new__(DocumentBuilder)
        sections = {"timelines_and_deliverables": "A source-grounded delivery sequence."}
        self.assertEqual(
            builder._resolve_content("3. Timeline and Deliverables", sections),
            "A source-grounded delivery sequence.",
        )

    def test_numbered_about_client_title_resolves_canonical_content_key(self):
        builder = DocumentBuilder.__new__(DocumentBuilder)
        metadata = {
            "company_name": "Example Customer",
            "author_org": "ShellKode",
        }
        self.assertEqual(
            builder._resolve_section_content(
                "2. About Example Customer",
                {"about_client": "A confirmed customer profile."},
                metadata,
            ),
            "A confirmed customer profile.",
        )

    def test_benchmark_order_chrome_and_wide_table_fallback(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            builder = DocumentBuilder(_Config(temp_dir))
            sections = {
                "toc_structure": (
                    "Document Version Control\n"
                    "1. Objective\n"
                    "Acceptance and Signatories to Statement of Work"
                ),
                "document_control_and_basis": (
                    "| Version | Date | Prepared By | Status | Classification |\n"
                    "|---|---|---|---|---|\n"
                    "| 1.0 | 18 August 2026 | Test Author | Draft for Client Review | Confidential |"
                ),
                "project_overview": (
                    "A source-grounded purpose.\n\n"
                    "| ID | Requirement | Source | Response | Output | Evidence | Owner | Status |\n"
                    "|---|---|---|---|---|---|---|---|\n"
                    "| R-1 | Route requests | BRD | Configure rules | Queue | Test log | Lead | Proposed |"
                ),
                "acceptance_and_signatories_to_statement_of_work": (
                    "| ShellKode | Example Customer |\n|---|---|\n"
                    "| Name: | Name: |\n| Title: | Title: |\n"
                    "| Signature: | Signature: |\n| Date: | Date: |"
                ),
            }
            metadata = {
                "company_name": "Example Customer",
                "project_title": "Agentic CRM Platform",
                "author_name": "Test Author",
                "author_org": "ShellKode",
                "document_date": "18 August 2026",
                "version": "1.0",
            }
            output = builder.build_document(sections, metadata, mode="POC")
            document = Document(output)
            body_text = [paragraph.text for paragraph in document.paragraphs]
            self.assertLess(body_text.index("Table of Contents"), body_text.index("Document Version Control"))
            self.assertLess(body_text.index("Document Version Control"), body_text.index("1. Objective"))
            self.assertEqual(len(document.tables), 4)

            signature_table = document.tables[-1]
            shading = signature_table.cell(0, 0)._tc.tcPr.find(qn("w:shd"))
            self.assertIsNotNone(shading)
            self.assertEqual(shading.get(qn("w:fill")), "5D3FD3")
            self.assertEqual(signature_table.cell(1, 0).text, "Name:")
            self.assertNotIn("Bhuvanesh", "\n".join(cell.text for row in signature_table.rows for cell in row.cells))

            control_header = document.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:shd"))
            self.assertEqual(control_header.get(qn("w:val")), "clear")
            self.assertEqual(control_header.get(qn("w:color")), "auto")
            self.assertEqual(control_header.get(qn("w:fill")), "5D3FD3")

            content_section = document.sections[-1]
            self.assertAlmostEqual(content_section.left_margin.inches, 0.68, places=2)
            self.assertIn("Confidential Copyright © ShellKode 2026", content_section.footer.paragraphs[0].text)
            header_table = content_section.header.tables[0]
            left_header = header_table.cell(0, 0).paragraphs[0]
            right_header = header_table.cell(0, 2).paragraphs[0]
            self.assertAlmostEqual(
                left_header.paragraph_format.left_indent.inches,
                CHROME_INSET_IN,
                places=2,
            )
            self.assertAlmostEqual(
                right_header.paragraph_format.right_indent.inches,
                CHROME_INSET_IN,
                places=2,
            )
            self.assertTrue(left_header.text.startswith("Agentic CRM Platform"))
            self.assertEqual(left_header.alignment, 0)
            self.assertEqual(right_header.alignment, 2)
            self.assertEqual(content_section.footer.paragraphs[1].alignment, 1)

            with zipfile.ZipFile(output) as package:
                document_xml = package.read("word/document.xml").decode("utf-8")
                header_xml = "\n".join(
                    package.read(name).decode("utf-8")
                    for name in package.namelist()
                    if re.fullmatch(r"word/header\d+\.xml", name)
                )
                footer_xml = "\n".join(
                    package.read(name).decode("utf-8")
                    for name in package.namelist()
                    if re.fullmatch(r"word/footer\d+\.xml", name)
                )
                settings_xml = package.read("word/settings.xml").decode("utf-8")
                styles_xml = package.read("word/styles.xml").decode("utf-8")
                font_table_xml = package.read("word/fontTable.xml").decode("utf-8")
                font_relationships_xml = package.read(
                    "word/_rels/fontTable.xml.rels"
                ).decode("utf-8")
                package_names = set(package.namelist())
            # The cover-to-TOC transition is a native next-page section break;
            # only the TOC-to-body transition needs an explicit page break.
            self.assertEqual(document_xml.count('<w:br w:type="page"'), 1)
            self.assertNotIn('<w:pageBreakBefore w:val="0"', document_xml)
            self.assertIn("w:drawing", header_xml)
            self.assertIn("w:tbl", header_xml)
            self.assertIn('w:gridCol w:w="7200"', header_xml)
            self.assertIn('w:jc w:val="left"', header_xml)
            self.assertIn('w:jc w:val="right"', header_xml)
            self.assertIn("PAGE", footer_xml)
            self.assertIn("NUMPAGES", footer_xml)
            expected_toc_tab = round(
                (PAGE_WIDTH_IN - 2 * HORIZONTAL_MARGIN_IN - 0.18) * 1440
            )
            self.assertIn(f'w:pos="{expected_toc_tab}"', document_xml)
            self.assertEqual(
                document_xml.count('TOC \\o "1-4" \\h \\z \\u'),
                1,
            )
            self.assertIn('<w:updateFields w:val="true"', settings_xml)
            self.assertNotIn("Courier", document_xml)
            for font_attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                self.assertIn(f'w:{font_attribute}="DM Sans"', styles_xml)
            self.assertIn('w:name="DM Sans"', font_table_xml)
            self.assertIn("w:embedRegular", font_table_xml)
            self.assertIn("w:embedBold", font_table_xml)
            self.assertIn("/relationships/font", font_relationships_xml)
            self.assertIn("word/fonts/DMSans-Regular.odttf", package_names)
            self.assertIn("word/fonts/DMSans-Bold.odttf", package_names)

            # The approved template keeps its decorative background and logo as
            # native floating artwork. Text remains in editable paragraphs.
            self.assertIn('behindDoc="1"', document_xml)
            self.assertIn('relativeFrom="page"', document_xml)
            self.assertIn("wp:anchor", document_xml)
            self.assertNotIn("Client Name", document_xml)
            self.assertNotIn("Project Name", document_xml)
            self.assertNotIn("&lt;Author Name&gt;", document_xml)
            self.assertNotIn("&lt;Today’s Date&gt;", document_xml)
            self.assertNotIn("&lt;Date&gt;", document_xml)
            self.assertNotIn("wp:inline", document_xml)
            self.assertIn("word/media/image1.png", package_names)

            # LibreOffice uses a separate layout engine from desktop Word and
            # catches the same class of section-boundary issue seen in Word
            # Online. The TOC must be the page immediately after the cover.
            office_binary = _find_libreoffice_binary(builder.config)
            if office_binary:
                profile_dir = Path(temp_dir) / "lo-profile"
                profile_dir.mkdir()
                completed = subprocess.run(
                    [
                        office_binary,
                        "--headless",
                        f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        temp_dir,
                        output,
                    ],
                    capture_output=True,
                    text=True,
                    timeout=90,
                    check=False,
                )
                self.assertEqual(completed.returncode, 0, completed.stderr)
                pdf = PdfReader(str(Path(output).with_suffix(".pdf")))
                self.assertIn("Example Customer", pdf.pages[0].extract_text())
                self.assertIn("Table of Contents", pdf.pages[1].extract_text())


if __name__ == "__main__":
    unittest.main()
