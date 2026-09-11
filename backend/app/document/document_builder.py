"""Professional, deterministic DOCX builder for generated SOW content.

Design system: the ShellKode benchmark proposal system using DM Sans,
benchmark-matched purple/blue accents, compact tables, and fixed Letter geometry.
First page pattern: ``proposal_centerpiece`` over the existing brand artwork.
"""

from __future__ import annotations

import base64
import copy
import io
import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from xml.sax.saxutils import escape as xml_escape
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from PIL import Image as PILImage, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

from app.diagram.service import (
    ASSET_KEY as ARCHITECTURE_ASSETS_KEY,
    LEGACY_ASSET_KEY as ARCHITECTURE_ASSET_KEY,
)


PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
HORIZONTAL_MARGIN_IN = 0.68
CHROME_INSET_IN = 0.08
TOP_MARGIN_IN = 0.66
BOTTOM_MARGIN_IN = 0.58
CONTENT_WIDTH_DXA = 10282
TABLE_INDENT_DXA = 0
FONT_NAME = "DM Sans"
PURPLE = "5D3FD3"
BLUE = "1A4BD2"
DARK_PURPLE = "3E2A91"
INK = "434343"
MUTED = "6F6F6F"
LIGHT_FILL = "F2F2F2"
BORDER = "D9D9D9"
EMBEDDED_FONT_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.obfuscatedFont"
FONT_RELATIONSHIP_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _set_font(run, name: str = FONT_NAME, size: Optional[float] = None,
              bold: Optional[bool] = None, italic: Optional[bool] = None,
              color: Optional[str] = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_character_spacing(element, twentieth_points: int = 2) -> None:
    """Apply subtle tracking without substituting or widening space glyphs."""
    rpr = element.get_or_add_rPr()
    spacing = rpr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        rpr.append(spacing)
    spacing.set(qn("w:val"), str(twentieth_points))


def _add_field(paragraph, instruction: str, cached_text: str = "", size: float = 9):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = cached_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    _set_font(run, size=size, color=MUTED)
    return run


def _start_complex_field(paragraph, instruction: str):
    """Start a multi-paragraph Word field before its cached display content."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, instr, separate])
    return run


def _end_complex_field(paragraph):
    """Close a complex field after the final cached-result paragraph."""
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)
    return run


def _bookmark_name(title: str, index: int = 0) -> str:
    prefix = f"SOW_{index}_"
    # Word bookmark names are limited to 40 characters. Keeping the stable,
    # unique numeric prefix and trimming only the descriptive slug prevents
    # LibreOffice from silently renaming the bookmark during pagination.
    slug_limit = max(0, 40 - len(prefix))
    slug = re.sub(r"[^A-Za-z0-9_]", "", title.replace(" ", "_"))[:slug_limit]
    return f"{prefix}{slug}" if slug else prefix.rstrip("_")


def _add_hyperlink(paragraph, text: str, anchor: str):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), INK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT_NAME)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "2")
    rpr.extend([fonts, color, underline, size, spacing])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_external_hyperlink(paragraph, text: str, url: str):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT_NAME)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "2")
    rpr.extend([fonts, color, underline, size, spacing])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _mark_update_fields(document: Document) -> None:
    """Ask Word renderers to resolve dynamic fields when no cache exists.

    When an optional pagination engine later provides cached PAGEREF values,
    the package flag is switched off so desktop Word preserves that cache.
    """
    settings = document.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def _set_update_fields_flag(document_path: Path, enabled: bool) -> None:
    """Set the package-wide field refresh flag without changing document layout."""
    with zipfile.ZipFile(document_path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}
    settings = parse_xml(files["word/settings.xml"])
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true" if enabled else "false")
    files["word/settings.xml"] = etree.tostring(
        settings, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{document_path.stem}-settings-", suffix=".docx", dir=document_path.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as destination:
            for name, payload in files.items():
                destination.writestr(name, payload)
        os.replace(temporary_path, document_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def _obfuscate_embedded_font(font_data: bytes, font_key: uuid.UUID) -> bytes:
    """Apply the ECMA-376 obfuscation used for embedded Word fonts."""
    data = bytearray(font_data)
    key = font_key.bytes
    for index in range(min(32, len(data))):
        data[index] ^= key[15 - (index % 16)]
    return bytes(data)


def _embed_dm_sans(document_path: Path, config) -> None:
    """Embed the bundled DM Sans regular/bold faces in a generated DOCX."""
    regular_path = Path(config.ASSETS_DIR) / "fonts" / "DMSans-Regular.ttf"
    bold_path = Path(config.ASSETS_DIR) / "fonts" / "DMSans-Bold.ttf"
    if not regular_path.exists() or not bold_path.exists():
        raise FileNotFoundError("Bundled DM Sans regular and bold font files are required")

    keys = {
        "regular": uuid.UUID("65c9e9d6-9297-5b37-a1a2-62186d121c7e"),
        "bold": uuid.UUID("c0bd415f-6064-5ea9-8cf0-72409c43db81"),
    }
    font_entries = {
        "word/fonts/DMSans-Regular.odttf": _obfuscate_embedded_font(regular_path.read_bytes(), keys["regular"]),
        "word/fonts/DMSans-Bold.odttf": _obfuscate_embedded_font(bold_path.read_bytes(), keys["bold"]),
    }

    with zipfile.ZipFile(document_path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}

    font_table = parse_xml(files["word/fontTable.xml"])
    for existing in list(font_table.findall(qn("w:font"))):
        if existing.get(qn("w:name")) == FONT_NAME:
            font_table.remove(existing)
    font = OxmlElement("w:font")
    font.set(qn("w:name"), FONT_NAME)
    family = OxmlElement("w:family")
    family.set(qn("w:val"), "swiss")
    pitch = OxmlElement("w:pitch")
    pitch.set(qn("w:val"), "variable")
    regular = OxmlElement("w:embedRegular")
    regular.set(qn("r:id"), "rIdDmSansRegular")
    regular.set(qn("w:fontKey"), "{" + str(keys["regular"]).upper() + "}")
    regular.set(qn("w:subsetted"), "false")
    bold = OxmlElement("w:embedBold")
    bold.set(qn("r:id"), "rIdDmSansBold")
    bold.set(qn("w:fontKey"), "{" + str(keys["bold"]).upper() + "}")
    bold.set(qn("w:subsetted"), "false")
    font.extend([family, pitch, regular, bold])
    font_table.append(font)
    files["word/fontTable.xml"] = etree.tostring(
        font_table, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    relationships_path = "word/_rels/fontTable.xml.rels"
    if relationships_path in files:
        relationships = parse_xml(files[relationships_path])
    else:
        relationships = etree.Element(
            f"{{{RELATIONSHIPS_NS}}}Relationships", nsmap={None: RELATIONSHIPS_NS}
        )
    embedded_relationship_ids = {"rIdDmSansRegular", "rIdDmSansBold"}
    for relationship in list(relationships):
        if relationship.get("Id") in embedded_relationship_ids:
            relationships.remove(relationship)
    for rel_id, target in (
        ("rIdDmSansRegular", "fonts/DMSans-Regular.odttf"),
        ("rIdDmSansBold", "fonts/DMSans-Bold.odttf"),
    ):
        relationship = etree.Element(f"{{{RELATIONSHIPS_NS}}}Relationship")
        relationship.set("Id", rel_id)
        relationship.set("Type", FONT_RELATIONSHIP_TYPE)
        relationship.set("Target", target)
        relationships.append(relationship)
    files[relationships_path] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    content_types = parse_xml(files["[Content_Types].xml"])
    has_odttf = any(
        child.get("Extension", "").casefold() == "odttf"
        for child in content_types
    )
    if not has_odttf:
        default = etree.Element(f"{{{CONTENT_TYPES_NS}}}Default")
        default.set("Extension", "odttf")
        default.set("ContentType", EMBEDDED_FONT_CONTENT_TYPE)
        content_types.append(default)
    files["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    files.update(font_entries)

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{document_path.stem}-", suffix=".docx", dir=document_path.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as destination:
            for name, payload in files.items():
                destination.writestr(name, payload)
        os.replace(temporary_path, document_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def _normalise_field_instruction(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _pageref_results_from_xml(document_xml: bytes) -> Dict[str, str]:
    """Return cached PAGEREF results keyed by their normalised instruction."""
    root = parse_xml(document_xml)
    results: Dict[str, str] = {}
    selector = ".//*[self::w:fldChar or self::w:instrText or self::w:t]"
    for paragraph in root.findall(".//" + qn("w:p")):
        instruction_parts: List[str] = []
        result_parts: List[str] = []
        in_field = False
        after_separator = False
        for node in paragraph.xpath(selector):
            if node.tag == qn("w:fldChar"):
                field_type = node.get(qn("w:fldCharType"))
                if field_type == "begin":
                    in_field = True
                    after_separator = False
                    instruction_parts = []
                    result_parts = []
                elif in_field and field_type == "separate":
                    after_separator = True
                elif in_field and field_type == "end":
                    instruction = _normalise_field_instruction("".join(instruction_parts))
                    if instruction.startswith("PAGEREF "):
                        results[instruction] = "".join(result_parts).strip()
                    in_field = False
                    after_separator = False
            elif in_field and node.tag == qn("w:instrText"):
                instruction_parts.append(node.text or "")
            elif in_field and after_separator and node.tag == qn("w:t"):
                result_parts.append(node.text or "")
    return results


def _patch_pageref_results(document_xml: bytes, results: Dict[str, str]) -> bytes:
    """Write computed PAGEREF values into the original OOXML without reflowing it."""
    root = parse_xml(document_xml)
    selector = ".//*[self::w:fldChar or self::w:instrText or self::w:t]"
    patched: set[str] = set()
    for paragraph in root.findall(".//" + qn("w:p")):
        nodes = paragraph.xpath(selector)
        instruction_parts: List[str] = []
        result_nodes: List[Any] = []
        begin_node = None
        in_field = False
        after_separator = False
        for node in nodes:
            if node.tag == qn("w:fldChar"):
                field_type = node.get(qn("w:fldCharType"))
                if field_type == "begin":
                    begin_node = node
                    in_field = True
                    after_separator = False
                    instruction_parts = []
                    result_nodes = []
                elif in_field and field_type == "separate":
                    after_separator = True
                elif in_field and field_type == "end":
                    instruction = _normalise_field_instruction("".join(instruction_parts))
                    if instruction in results and result_nodes:
                        result_nodes[0].text = results[instruction]
                        for extra in result_nodes[1:]:
                            extra.text = ""
                        if begin_node is not None:
                            begin_node.set(qn("w:dirty"), "false")
                        patched.add(instruction)
                    in_field = False
                    after_separator = False
            elif in_field and node.tag == qn("w:instrText"):
                instruction_parts.append(node.text or "")
            elif in_field and after_separator and node.tag == qn("w:t"):
                result_nodes.append(node)

    missing = set(results) - patched
    if missing:
        raise RuntimeError(
            "Could not cache page numbers for PAGEREF field(s): "
            + ", ".join(sorted(missing))
        )
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _flatten_pageref_fields(document_xml: bytes) -> bytes:
    """Materialise cached PAGEREF values as text for Google Docs imports.

    Google Docs commonly ignores Word's complex PAGEREF result even when its
    cached ``w:t`` value is correct. Once an Office layout engine has calculated
    pagination, replace only PAGEREF field runs with ordinary text. Heading
    hyperlinks/bookmarks remain intact, and Word displays the same stable value.
    """
    root = parse_xml(document_xml)
    flattened = 0
    for paragraph in root.findall(".//" + qn("w:p")):
        children = list(paragraph)
        ranges: List[Tuple[int, int, str, Optional[Any]]] = []
        in_field = False
        after_separator = False
        start_index = -1
        instruction_parts: List[str] = []
        result_parts: List[str] = []
        result_properties = None
        for index, child in enumerate(children):
            nodes = child.xpath(".//*[self::w:fldChar or self::w:instrText or self::w:t]")
            for node in nodes:
                if node.tag == qn("w:fldChar"):
                    field_type = node.get(qn("w:fldCharType"))
                    if field_type == "begin":
                        in_field = True
                        after_separator = False
                        start_index = index
                        instruction_parts = []
                        result_parts = []
                        result_properties = None
                    elif in_field and field_type == "separate":
                        after_separator = True
                    elif in_field and field_type == "end":
                        instruction = _normalise_field_instruction("".join(instruction_parts))
                        value = "".join(result_parts).strip()
                        if instruction.startswith("PAGEREF ") and value.isdigit():
                            ranges.append((start_index, index, value, result_properties))
                        in_field = False
                        after_separator = False
                elif in_field and node.tag == qn("w:instrText"):
                    instruction_parts.append(node.text or "")
                elif in_field and after_separator and node.tag == qn("w:t"):
                    result_parts.append(node.text or "")
                    if result_properties is None:
                        run = node.getparent()
                        if run is not None and run.tag == qn("w:r"):
                            rpr = run.find(qn("w:rPr"))
                            if rpr is not None:
                                result_properties = copy.deepcopy(rpr)
        for start, end, value, properties in reversed(ranges):
            run = OxmlElement("w:r")
            if properties is not None:
                run.append(properties)
            text_node = OxmlElement("w:t")
            text_node.text = value
            run.append(text_node)
            for child in children[start:end + 1]:
                paragraph.remove(child)
            paragraph.insert(start, run)
            flattened += 1
    if not flattened:
        raise RuntimeError("No calculated PAGEREF fields were available to materialise")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _find_libreoffice_binary(config) -> Optional[str]:
    """Locate LibreOffice even when the backend was started with a minimal PATH."""
    configured = getattr(config, "LIBREOFFICE_BINARY", None) or os.getenv(
        "LIBREOFFICE_BINARY"
    )
    program_files = Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
    program_files_x86 = Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"))
    candidates = [
        Path(str(configured)).expanduser() if configured else None,
        Path(shutil.which("soffice")) if shutil.which("soffice") else None,
        Path(shutil.which("libreoffice")) if shutil.which("libreoffice") else None,
        program_files / "LibreOffice/program/soffice.exe",
        program_files_x86 / "LibreOffice/program/soffice.exe",
        Path.home()
        / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice",
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path("/opt/homebrew/bin/soffice"),
        Path("/usr/local/bin/soffice"),
        Path("/usr/bin/soffice"),
    ]
    for candidate in candidates:
        if candidate and candidate.is_file() and os.access(candidate, os.X_OK):
            return str(candidate)
    return None


def _refresh_fields_with_word(document_path: Path) -> None:
    """Use an existing desktop Word installation to populate cached fields.

    This has no Python or paid runtime dependency: it is an optional Windows
    optimisation when Word is already installed on the machine producing the
    document. The generated DOCX remains valid when Word is unavailable.
    """
    if os.name != "nt":
        raise RuntimeError("Microsoft Word field refresh is available only on Windows")
    script = r"""
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($env:SOW_DOCX_PATH, $false, $false)
    $document.Repaginate()
    [void]$document.Fields.Update()
    foreach ($toc in $document.TablesOfContents) { [void]$toc.Update() }
    foreach ($storyType in $document.StoryRanges) {
        $story = $storyType
        while ($null -ne $story) {
            [void]$story.Fields.Update()
            $story = $story.NextStoryRange
        }
    }
    $document.Repaginate()
    [void]$document.Fields.Update()
    $document.Save()
} finally {
    if ($null -ne $document) { $document.Close($false) }
    if ($null -ne $word) { $word.Quit() }
}
"""
    environment = os.environ.copy()
    environment["SOW_DOCX_PATH"] = str(document_path.resolve())
    completed = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
        env=environment,
    )
    if completed.returncode != 0:
        details = (completed.stderr or completed.stdout or "Word COM is unavailable").strip()
        raise RuntimeError(f"Microsoft Word could not refresh document fields: {details}")
    with zipfile.ZipFile(document_path, "r") as package:
        results = _pageref_results_from_xml(package.read("word/document.xml"))
    if not results or any(not value.isdigit() for value in results.values()):
        raise RuntimeError("Microsoft Word did not return numeric TOC page references")
    _set_update_fields_flag(document_path, False)


def _refresh_pageref_cached_results(document_path: Path, config) -> None:
    """Refresh the native TOC with an available Office layout engine.

    Existing Microsoft Word is preferred on Windows. LibreOffice remains a
    free, optional fallback; neither application is a hard project dependency.
    The native TOC and its nested PAGEREF fields must remain intact so Word and
    Google Docs can rebuild entry text after a heading is renamed.
    """
    if not getattr(config, "PRECOMPUTE_DOCUMENT_FIELDS", True):
        raise RuntimeError("optional field-cache generation is disabled")

    word_error: Optional[Exception] = None
    if os.name == "nt":
        try:
            _refresh_fields_with_word(document_path)
            return
        except Exception as exc:
            word_error = exc

    office_binary = _find_libreoffice_binary(config)
    if not office_binary:
        detail = f"; Word refresh failed: {word_error}" if word_error else ""
        raise RuntimeError("no optional document pagination engine is available" + detail)

    with tempfile.TemporaryDirectory(
        prefix=f".{document_path.stem}-fields-", dir=document_path.parent
    ) as temp_name:
        temp_root = Path(temp_name)
        input_dir = temp_root / "input"
        output_dir = temp_root / "output"
        profile_dir = temp_root / "profile"
        input_dir.mkdir()
        output_dir.mkdir()
        profile_dir.mkdir()
        source_copy = input_dir / document_path.name
        shutil.copy2(document_path, source_copy)
        command = [
            str(office_binary),
            "--headless",
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--convert-to",
            "docx:Office Open XML Text",
            "--outdir",
            str(output_dir),
            str(source_copy),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        refreshed_path = output_dir / document_path.name
        if completed.returncode != 0 or not refreshed_path.exists():
            details = (completed.stderr or completed.stdout or "unknown error").strip()
            raise RuntimeError(f"Failed to calculate TOC page numbers: {details}")

        with zipfile.ZipFile(refreshed_path, "r") as refreshed_package:
            refreshed_results = _pageref_results_from_xml(
                refreshed_package.read("word/document.xml")
            )
        if not refreshed_results or any(not value.isdigit() for value in refreshed_results.values()):
            raise RuntimeError("LibreOffice did not return numeric TOC page references")

        with zipfile.ZipFile(document_path, "r") as source_package:
            files = {name: source_package.read(name) for name in source_package.namelist()}
        original_results = _pageref_results_from_xml(files["word/document.xml"])
        missing_results = set(original_results) - set(refreshed_results)
        if missing_results:
            raise RuntimeError(
                "Pagination did not resolve every TOC entry: "
                + ", ".join(sorted(missing_results))
            )
        files["word/document.xml"] = _patch_pageref_results(
            files["word/document.xml"], refreshed_results
        )
        settings_root = parse_xml(files["word/settings.xml"])
        update_fields = settings_root.find(qn("w:updateFields"))
        if update_fields is None:
            update_fields = OxmlElement("w:updateFields")
            settings_root.append(update_fields)
        update_fields.set(qn("w:val"), "false")
        files["word/settings.xml"] = etree.tostring(
            settings_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{document_path.stem}-cached-",
            suffix=".docx",
            dir=document_path.parent,
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as destination:
                for name, payload in files.items():
                    destination.writestr(name, payload)
            os.replace(temporary_path, document_path)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()


class SectionBuilder:
    """Render ordered Markdown blocks without losing text after tables."""

    def __init__(self, doc: Document, config):
        self.doc = doc
        self.config = config
        self.bullet_num_id = 91

    def parse_content(
        self,
        content: str,
        heading_prefix: Optional[str] = None,
        heading_anchor_map: Optional[Dict[str, str]] = None,
        bookmark_callback: Optional[Any] = None,
        section_title: Optional[str] = None,
        table_caption_callback: Optional[Any] = None,
        heading_callback: Optional[Any] = None,
    ) -> List[Any]:
        added: List[Any] = []
        caption_context = section_title or "Statement of Work"
        active_heading: Optional[Tuple[str, int]] = None
        heading_counters = {2: 0, 3: 0, 4: 0}
        for kind, payload in self._iter_blocks(content or ""):
            if kind == "table":
                tables = self.build_table(payload)
                if tables is not None:
                    table_items = tables if isinstance(tables, list) else [tables]
                    for table_index, table in enumerate(table_items, 1):
                        added.append(table)
                        spacing_anchor = table._tbl
                        if table_caption_callback:
                            context = caption_context
                            if len(table_items) > 1:
                                context = f"{context} (continued {table_index})"
                            caption = table_caption_callback(table, context)
                            if caption is not None:
                                # build_table may create continuation tables in
                                # one pass. Re-anchor each caption immediately
                                # after the table it describes.
                                table._tbl.addnext(caption._p)
                                added.append(caption)
                                spacing_anchor = caption._p
                        spacer = self.doc.add_paragraph()
                        spacer.paragraph_format.space_before = Pt(0)
                        spacer.paragraph_format.space_after = Pt(5)
                        spacer.paragraph_format.line_spacing = Pt(1)
                        spacing_anchor.addnext(spacer._p)
                        added.append(spacer)
            elif kind == "heading":
                if active_heading and heading_callback:
                    inserted = heading_callback(*active_heading)
                    if inserted:
                        added.extend(inserted if isinstance(inserted, list) else [inserted])
                level, text = payload
                text = self._numbered_heading_text(
                    text, level, heading_prefix, heading_counters
                )
                paragraph = self.doc.add_paragraph(style=f"Heading {level}")
                self._add_rich_runs(paragraph, text, size=15.0)
                anchor = (heading_anchor_map or {}).get(text)
                if anchor and bookmark_callback:
                    bookmark_callback(paragraph, anchor)
                added.append(paragraph)
                caption_context = re.sub(r"^\d+(?:\.\d+)*\s+", "", text).strip() or caption_context
                active_heading = (text, level)
            elif kind == "bullet":
                level, text = payload
                paragraph = self.doc.add_paragraph()
                self._add_rich_runs(paragraph, text)
                self._apply_numbering(paragraph, self.bullet_num_id, level)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                added.append(paragraph)
            elif kind == "number":
                level, text = payload
                paragraph = self.doc.add_paragraph()
                self._add_rich_runs(paragraph, text)
                # The deliverable uses bullets exclusively. Converting ordered
                # Markdown here also prevents Word from continuing a hidden
                # decimal sequence across later sections (for example at 59).
                self._apply_numbering(paragraph, self.bullet_num_id, level)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                added.append(paragraph)
            elif kind == "caption":
                paragraph = self.doc.add_paragraph(style="Caption")
                self._add_rich_runs(paragraph, payload)
                added.append(paragraph)
            else:
                paragraph = self.doc.add_paragraph()
                self._add_rich_runs(paragraph, payload)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                added.append(paragraph)
        if active_heading and heading_callback:
            inserted = heading_callback(*active_heading)
            if inserted:
                added.extend(inserted if isinstance(inserted, list) else [inserted])
        return added

    @classmethod
    def enumerate_headings(
        cls, content: str, heading_prefix: Optional[str]
    ) -> List[Tuple[str, int]]:
        """Return the same numbered heading labels used during rendering."""
        counters = {2: 0, 3: 0, 4: 0}
        result: List[Tuple[str, int]] = []
        for kind, payload in cls._iter_blocks(content or ""):
            if kind != "heading":
                continue
            level, text = payload
            result.append((cls._numbered_heading_text(text, level, heading_prefix, counters), level))
        return result

    @staticmethod
    def _numbered_heading_text(
        text: str,
        level: int,
        heading_prefix: Optional[str],
        counters: Dict[int, int],
    ) -> str:
        clean = re.sub(r"^\d+(?:\.\d+)*[.)]?\s*", "", str(text)).strip()
        if not heading_prefix or not re.fullmatch(r"\d+", str(heading_prefix)):
            return clean
        level = min(4, max(2, level))
        counters[level] += 1
        for child_level in range(level + 1, 5):
            counters[child_level] = 0
        # A malformed H3/H4 without its parent should still receive a stable,
        # non-zero hierarchy rather than displaying e.g. 4.0.1.
        for parent_level in range(2, level):
            if counters[parent_level] == 0:
                counters[parent_level] = 1
        suffix = ".".join(str(counters[item]) for item in range(2, level + 1))
        return f"{heading_prefix}.{suffix} {clean}"

    @staticmethod
    def _iter_blocks(content: str):
        lines = content.replace("<br/>", "<br>").splitlines()
        index = 0
        prose: List[str] = []

        def flush_prose():
            nonlocal prose
            if prose:
                joined = " ".join(part.strip() for part in prose if part.strip())
                prose = []
                if joined:
                    return ("prose", joined)
            return None

        while index < len(lines):
            raw = lines[index]
            stripped = raw.strip()
            if not stripped:
                block = flush_prose()
                if block:
                    yield block
                index += 1
                continue
            if re.match(r"^\s*\|.+\|\s*$", raw):
                block = flush_prose()
                if block:
                    yield block
                table_lines = []
                while index < len(lines) and re.match(r"^\s*\|.+\|\s*$", lines[index]):
                    table_lines.append(lines[index])
                    index += 1
                yield ("table", table_lines)
                continue
            heading = re.match(r"^(#{3,5})\s+(.+)$", stripped)
            if heading:
                block = flush_prose()
                if block:
                    yield block
                level = min(4, max(2, len(heading.group(1)) - 1))
                yield ("heading", (level, heading.group(2).strip()))
                index += 1
                continue
            bullet = re.match(r"^(\s*)[-*•○]\s+(.+)$", raw)
            if bullet:
                block = flush_prose()
                if block:
                    yield block
                yield ("bullet", (min(len(bullet.group(1)) // 2, 2), bullet.group(2).strip()))
                index += 1
                continue
            numbered = re.match(r"^(\s*)\d+[.)]\s+(.+)$", raw)
            if numbered:
                block = flush_prose()
                if block:
                    yield block
                yield ("number", (min(len(numbered.group(1)) // 2, 2), numbered.group(2).strip()))
                index += 1
                continue
            if re.match(r"^(Table|Figure)\s+\d*[:.-]", stripped, re.I):
                block = flush_prose()
                if block:
                    yield block
                # Captions are intentionally excluded from generated SOWs.
                index += 1
                continue
            if stripped.startswith("#"):
                index += 1
                continue
            prose.append(stripped)
            index += 1
        block = flush_prose()
        if block:
            yield block

    def _add_rich_runs(self, paragraph, text: str, size: Optional[float] = None, color: Optional[str] = None):
        parts = re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`.+?`)", str(text))
        for part in parts:
            if not part:
                continue
            bold = part.startswith("**") and part.endswith("**")
            italic = not bold and part.startswith("*") and part.endswith("*")
            code = part.startswith("`") and part.endswith("`")
            clean = part[2:-2] if bold else part[1:-1] if italic or code else part
            segments = re.split(r"(https?://[^\s<>()\[\]]+)", clean)
            for segment in segments:
                if not segment:
                    continue
                if re.fullmatch(r"https?://[^\s<>()\[\]]+", segment):
                    url = segment.rstrip(".,;:")
                    _add_external_hyperlink(paragraph, url, url)
                    trailing = segment[len(url):]
                    if trailing:
                        run = paragraph.add_run(trailing)
                        _set_font(run, FONT_NAME, size=size, bold=bold, italic=italic, color=color)
                    continue
                run = paragraph.add_run(segment)
                # Keep the document type system consistent even for inline code.
                # Backticks still distinguish the source semantically; typography
                # remains DM Sans as required by the SOW brand standard.
                _set_font(run, FONT_NAME, size=size, bold=bold, italic=italic, color=color)

    @staticmethod
    def _apply_numbering(paragraph, num_id: int, level: int) -> None:
        ppr = paragraph._element.get_or_add_pPr()
        numpr = ppr.find(qn("w:numPr"))
        if numpr is not None:
            ppr.remove(numpr)
        numpr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        numid = OxmlElement("w:numId")
        numid.set(qn("w:val"), str(num_id))
        numpr.extend([ilvl, numid])
        ppr.append(numpr)

    @staticmethod
    def _parse_table_rows(lines: Sequence[str]) -> List[List[str]]:
        rows: List[List[str]] = []
        for line in lines:
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
                continue
            if cells:
                rows.append(cells)
        if not rows:
            return []
        width = max(len(row) for row in rows)
        return [row + [""] * (width - len(row)) for row in rows]

    @staticmethod
    def _column_widths(rows: Sequence[Sequence[str]]) -> List[int]:
        columns = len(rows[0])
        headers = [cell.lower() for cell in rows[0]]
        # Stable semantic layouts keep control fields compact and reserve room
        # for the narrative evidence/criteria columns that need it most.
        semantic_patterns = (
            (("module", "open item", "status"), [2450, 5000, 2832]),
            (("area", "open item", "note"), [2450, 5000, 2832]),
            (("week", "phase", "activities", "deliverable", "dependency", "owner"),
             [720, 1200, 2700, 1800, 1500, 1440]),
            (("timeframe", "phase", "activities", "deliverable", "customer", "exit"),
             [900, 1400, 2450, 1600, 1400, 1610]),
            (("id", "requirement", "priority", "description", "source", "acceptance"),
             [650, 1750, 820, 2050, 1400, 2690]),
            (("id", "requirement", "status", "delivery", "output", "validation"),
             [650, 2200, 900, 1900, 1100, 2610]),
            (("id", "deliverable", "format", "owner", "timing", "acceptance"),
             [650, 2300, 1050, 1100, 1300, 2960]),
            (("id", "deliverable", "format", "owner", "evidence", "timing"),
             [650, 2400, 1100, 1000, 2800, 1410]),
            (("component", "purpose", "data", "integration", "security", "owner"),
             [1400, 1800, 1500, 1500, 1900, 1260]),
            (("layer", "component", "purpose", "data", "control", "status"),
             [1200, 1800, 1800, 1500, 1800, 1260]),
            (("id", "risk", "likelihood", "mitigation", "contingency", "owner"),
             [650, 1800, 1200, 2300, 2200, 1210]),
            (("id", "risk", "likelihood", "impact", "mitigation", "contingency", "owner", "trigger"),
             [550, 1500, 850, 700, 1650, 1550, 1000, 1560]),
            (("id", "risk", "cause", "likelihood", "impact", "mitigation", "contingency", "owner"),
             [540, 1200, 1050, 850, 750, 1800, 1800, 1370]),
        )
        for signature, widths in semantic_patterns:
            if columns == len(signature) and all(token in headers[index] for index, token in enumerate(signature)):
                scaled = [round(width * CONTENT_WIDTH_DXA / sum(widths)) for width in widths]
                scaled[-1] += CONTENT_WIDTH_DXA - sum(scaled)
                return scaled
        patterns = {
            2: {
                ("timeframe", "milestones"): [2100, 7260],
                ("id", "description"): [1500, 7860],
            },
            3: {},
        }
        for signature, widths in patterns.get(columns, {}).items():
            if all(any(token in headers[i] for token in (signature[i],)) for i in range(columns)):
                scaled = [round(width * CONTENT_WIDTH_DXA / sum(widths)) for width in widths]
                scaled[-1] += CONTENT_WIDTH_DXA - sum(scaled)
                return scaled
        scores = []
        for col in range(columns):
            lengths = [min(len(re.sub(r"<br>", " ", row[col])), 120) for row in rows]
            header = headers[col]
            score = max(10, sum(lengths) / max(1, len(lengths)))
            if any(token in header for token in ("id", "week", "status", "date", "owner", "role", "priority")):
                score *= 0.65
            if any(token in header for token in ("description", "activities", "deliverable", "evidence", "criteria", "requirement", "mitigation")):
                score *= 1.5
            scores.append(max(score, 8))
        minimum = 1050 if columns <= 6 else 800
        widths = [max(minimum, round(CONTENT_WIDTH_DXA * score / sum(scores))) for score in scores]
        while sum(widths) > CONTENT_WIDTH_DXA:
            index = max(range(columns), key=lambda i: widths[i])
            if widths[index] <= minimum:
                break
            widths[index] -= 10
        while sum(widths) < CONTENT_WIDTH_DXA:
            widths[max(range(columns), key=lambda i: scores[i])] += 1
        return widths

    def build_table(self, lines: Sequence[str]):
        rows = self._parse_table_rows(lines)
        if not rows:
            return None
        if len(rows[0]) > 5:
            # Word cannot keep six-to-eight narrative columns readable on a
            # portrait Letter page. Preserve every value while splitting the
            # record into continuation tables that repeat the identifier.
            tables = []
            for start in range(1, len(rows[0]), 4):
                columns = [0] + list(range(start, min(start + 4, len(rows[0]))))
                continuation_rows = [[row[index] for index in columns] for row in rows]
                tables.append(self._build_table_rows(continuation_rows))
            return tables
        return self._build_table_rows(rows)

    def _build_table_rows(self, rows: Sequence[Sequence[str]]):
        cols = len(rows[0])
        # A single-row Markdown table is used by the legacy signature block
        # for signer names and acceptance dates. It is a continuation/data
        # row, not a header, so do not render it as a purple header band.
        has_header = len(rows) > 1
        widths = self._column_widths(rows)
        font_size = 11.0
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.style = "SOW Table"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        self._set_table_geometry(table, widths)
        self._set_table_borders(table)

        for row_index, values in enumerate(rows):
            row = table.rows[row_index]
            self._set_row_cant_split(row)
            if has_header and row_index == 0:
                self._set_repeat_header(row)
            for col_index, value in enumerate(values):
                cell = row.cells[col_index]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self._set_cell_margins(cell)
                self._set_cell_width(cell, widths[col_index])
                fill = PURPLE if has_header and row_index == 0 else (LIGHT_FILL if row_index % 2 == 1 else "FFFFFF")
                self._set_cell_shading(cell, fill)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pieces = re.split(r"<br\s*/?>", value, flags=re.I)
                for piece_index, piece in enumerate(pieces):
                    if piece_index:
                        paragraph.add_run().add_break(WD_BREAK.LINE)
                    self._add_rich_runs(
                        paragraph,
                        re.sub(r"^[•-]\s*", "", piece.strip()),
                        size=font_size,
                        color="FFFFFF" if has_header and row_index == 0 else INK,
                    )
                for run in paragraph.runs:
                    if has_header and row_index == 0:
                        run.bold = True
        return table

    @staticmethod
    def _set_table_geometry(table, widths: Sequence[int]) -> None:
        tbl = table._tbl
        tblpr = tbl.tblPr
        layout = tblpr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tblpr.append(layout)
        layout.set(qn("w:type"), "fixed")
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
        tblw.set(qn("w:type"), "dxa")
        indent = tblpr.find(qn("w:tblInd"))
        if indent is None:
            indent = OxmlElement("w:tblInd")
            tblpr.append(indent)
        indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
        indent.set(qn("w:type"), "dxa")
        grid = tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(width))
            grid.append(col)

    @staticmethod
    def _set_table_borders(table) -> None:
        tblpr = table._tbl.tblPr
        borders = tblpr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblpr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = borders.find(qn(f"w:{edge}"))
            if element is None:
                element = OxmlElement(f"w:{edge}")
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), BORDER)

    @staticmethod
    def _set_cell_width(cell, width: int) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        tcw = tcpr.find(qn("w:tcW"))
        if tcw is None:
            tcw = OxmlElement("w:tcW")
            tcpr.append(tcw)
        tcw.set(qn("w:w"), str(width))
        tcw.set(qn("w:type"), "dxa")

    @staticmethod
    def _set_cell_margins(cell) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        margins = tcpr.find(qn("w:tcMar"))
        if margins is None:
            margins = OxmlElement("w:tcMar")
            tcpr.append(margins)
        for side, value in (("top", 55), ("bottom", 55), ("start", 75), ("end", 75)):
            element = margins.find(qn(f"w:{side}"))
            if element is None:
                element = OxmlElement(f"w:{side}")
                margins.append(element)
            element.set(qn("w:w"), str(value))
            element.set(qn("w:type"), "dxa")

    @staticmethod
    def _set_cell_shading(cell, fill: str) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        shading = tcpr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tcpr.append(shading)
        # Word requires an explicit shading pattern; LibreOffice is more
        # forgiving. Without w:val, Word may omit the purple fill while the
        # white header text remains, making every table header look blank.
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)

    @staticmethod
    def _set_repeat_header(row) -> None:
        trpr = row._tr.get_or_add_trPr()
        header = trpr.find(qn("w:tblHeader"))
        if header is None:
            header = OxmlElement("w:tblHeader")
            trpr.append(header)
        header.set(qn("w:val"), "1")

    @staticmethod
    def _set_row_cant_split(row) -> None:
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))

    def has_table(self, content: str) -> bool:
        return bool(re.search(r"(?m)^\s*\|.+\|\s*$", content or ""))

    def parse_table(self, content: str):
        """Compatibility helper that retains all non-table text in source order."""
        table_lines = [line for line in (content or "").splitlines() if re.match(r"^\s*\|.+\|\s*$", line)]
        other = [line for line in (content or "").splitlines() if line not in table_lines and line.strip()]
        return other, table_lines


class DocumentBuilder:
    """Build a branded, accessible, navigable DOCX from section Markdown."""

    RESERVED_KEYS = {
        "cover_page", "toc_structure", "table_of_contents", "tableofcontents",
        "table_contents", "generation_quality_summary", ARCHITECTURE_ASSET_KEY,
        ARCHITECTURE_ASSETS_KEY,
    }

    def __init__(self, config):
        self.config = config
        self.doc: Optional[Document] = None
        self.section_builder: Optional[SectionBuilder] = None
        self.toc_entries: List[str] = []
        self.expanded_toc_entries: List[Tuple[str, str, int]] = []
        self.subheading_anchor_maps: Dict[str, Dict[str, str]] = {}
        self._bookmark_ids = 0
        self._figure_counter = 0
        self._table_counter = 0

    @staticmethod
    def _get_short_company_name(company_name: str) -> str:
        name = company_name or ""
        for suffix in (" Pvt Ltd", " Private Limited", " Ltd", " LLC", " Inc", " Corporation", " Corp", " Limited", " Company"):
            if name.lower().endswith(suffix.lower()):
                return name[:-len(suffix)].strip()
        return name

    @staticmethod
    def _name_to_key(name: str) -> str:
        name = re.sub(r"^\s*\d+(?:\.\d+)*[.)]?\s*", "", name)
        return re.sub(r"[-\s]+", "_", re.sub(r"[^\w\s-]", "", name.lower())).strip("_")

    @staticmethod
    def _is_similar(section_name: str, key: str) -> bool:
        stop = {"the", "of", "and", "to", "a", "in", "for"}
        words_a = set(re.findall(r"\w+", section_name.lower())) - stop
        words_b = set(re.findall(r"\w+", key.replace("_", " ").lower())) - stop
        return bool(words_a and words_b and len(words_a & words_b) / max(len(words_a), len(words_b)) >= 0.55)

    def _parse_toc_entries(self, sections: Dict[str, Any]) -> None:
        content = next((str(sections[key]) for key in ("toc_structure", "table_of_contents", "tableofcontents", "table_contents") if sections.get(key)), "")
        entries: List[str] = []
        lines = [line for line in content.splitlines() if line.strip()]
        for line in lines:
            match = re.match(r"^\s*\d+[.)]\s+(.+?)\s*$", line)
            if match:
                # Explicit template numbering is part of the document contract;
                # never discard it based on the presence/absence of front matter.
                entries.append(line.strip())
                continue
            plain = line.strip()
            if plain and not plain.startswith(("#", "|", "[")):
                entries.append(plain)
        if not entries:
            entries = [key.replace("_", " ").title() for key, value in sections.items() if key not in self.RESERVED_KEYS and value]
        self.toc_entries = entries

    def _configure_styles(self) -> None:
        assert self.doc is not None
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = FONT_NAME
        normal.font.size = Pt(11)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            normal._element.rPr.rFonts.set(qn(f"w:{attr}"), FONT_NAME)
        normal.paragraph_format.space_before = Pt(0)
        _set_character_spacing(normal._element, 2)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.25
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.widow_control = True

        tokens = {
            "Heading 1": (20.0, PURPLE, 18, 8),
            "Heading 2": (15.0, BLUE, 12, 5),
            "Heading 3": (15.0, DARK_PURPLE, 9, 4),
            "Heading 4": (15.0, INK, 7, 3),
        }
        for style_name, (size, color, before, after) in tokens.items():
            style = styles[style_name]
            style.font.name = FONT_NAME
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(color)
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                style._element.rPr.rFonts.set(qn(f"w:{attr}"), FONT_NAME)
            _set_character_spacing(style._element, 2)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.widow_control = True
            style.paragraph_format.page_break_before = False
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Word Online applies named table styles more consistently than direct
        # cell formatting. Keep both layers: the style is the interoperable
        # contract and direct properties are a fallback for other renderers.
        table_style_name = "SOW Table"
        if table_style_name in styles:
            table_style = styles[table_style_name]
        else:
            table_style = styles.add_style(table_style_name, WD_STYLE_TYPE.TABLE)
        style_element = table_style._element
        for child_name in ("w:basedOn", "w:uiPriority", "w:tblPr", "w:tblStylePr"):
            for child in list(style_element.findall(qn(child_name))):
                style_element.remove(child)
        based_on = OxmlElement("w:basedOn")
        based_on.set(qn("w:val"), "TableGrid")
        priority = OxmlElement("w:uiPriority")
        priority.set(qn("w:val"), "40")
        table_properties = OxmlElement("w:tblPr")
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), BORDER)
            borders.append(border)
        table_properties.append(borders)

        first_row = OxmlElement("w:tblStylePr")
        first_row.set(qn("w:type"), "firstRow")
        run_properties = OxmlElement("w:rPr")
        bold = OxmlElement("w:b")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "FFFFFF")
        fonts = OxmlElement("w:rFonts")
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), FONT_NAME)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "22")
        run_properties.extend([fonts, bold, color, size])
        cell_properties = OxmlElement("w:tcPr")
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), PURPLE)
        cell_properties.append(shading)
        first_row.extend([run_properties, cell_properties])
        style_element.extend([based_on, priority, table_properties, first_row])

    def _add_numbering(self) -> None:
        assert self.doc is not None
        numbering = self.doc.part.numbering_part.element
        for abstract_id, num_id, fmt, text_values in (
            (91, 91, "bullet", ["•", "○", "▪"]),
        ):
            abstract = OxmlElement("w:abstractNum")
            abstract.set(qn("w:abstractNumId"), str(abstract_id))
            multi = OxmlElement("w:multiLevelType")
            multi.set(qn("w:val"), "hybridMultilevel")
            abstract.append(multi)
            for level in range(3):
                lvl = OxmlElement("w:lvl")
                lvl.set(qn("w:ilvl"), str(level))
                start = OxmlElement("w:start")
                start.set(qn("w:val"), "1")
                numfmt = OxmlElement("w:numFmt")
                numfmt.set(qn("w:val"), fmt)
                lvltext = OxmlElement("w:lvlText")
                lvltext.set(qn("w:val"), text_values[level])
                lvl_just = OxmlElement("w:lvlJc")
                lvl_just.set(qn("w:val"), "left")
                ppr = OxmlElement("w:pPr")
                tabs = OxmlElement("w:tabs")
                tab = OxmlElement("w:tab")
                tab.set(qn("w:val"), "num")
                tab.set(qn("w:pos"), str(540 + level * 360))
                tabs.append(tab)
                indent = OxmlElement("w:ind")
                indent.set(qn("w:left"), str(540 + level * 360))
                indent.set(qn("w:hanging"), "280")
                spacing = OxmlElement("w:spacing")
                spacing.set(qn("w:after"), "80")
                spacing.set(qn("w:line"), "290")
                spacing.set(qn("w:lineRule"), "auto")
                ppr.extend([tabs, indent, spacing])
                lvl.extend([start, numfmt, lvltext, lvl_just, ppr])
                abstract.append(lvl)
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(num_id))
            ref = OxmlElement("w:abstractNumId")
            ref.set(qn("w:val"), str(abstract_id))
            num.append(ref)
            numbering.extend([abstract, num])

    def build_document(self, sections: Dict[str, Any], metadata: Dict[str, Any], mode: str = "POC") -> Optional[Path]:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        prefix = "PROD" if mode in ("PROD", "POC_TO_PROD") else "POC"
        safe_company = re.sub(r"[^A-Za-z0-9_-]+", "_", metadata.get("company_name", "Client")).strip("_")
        path = Path(self.config.OUTPUT_DIR) / f"{prefix}_{safe_company}_{timestamp}.docx"
        path.parent.mkdir(parents=True, exist_ok=True)
        self._parse_toc_entries(sections)
        try:
            self._build_docx(path, sections, metadata, mode)
            return path
        except Exception as exc:
            print(f"❌ Error building DOCX: {exc}")
            import traceback
            traceback.print_exc()
            return None

    def _cover_template_path(self) -> Path:
        configured = getattr(self.config, "COVER_PAGE_TEMPLATE", None)
        if configured:
            return Path(configured)
        templates_dir = getattr(self.config, "TEMPLATES_DIR", None)
        if templates_dir:
            return Path(templates_dir) / "sow_coverpage_template.docx"
        return Path(self.config.ASSETS_DIR).parent / "templates" / "sow_coverpage_template.docx"

    @staticmethod
    def _fit_cover_text(draw, text: str, font_path: Path, preferred_size: int,
                        max_width: int, minimum_size: int = 22):
        size = preferred_size
        while size > minimum_size:
            font = ImageFont.truetype(str(font_path), size)
            if draw.textbbox((0, 0), text, font=font)[2] <= max_width:
                return font
            size -= 2
        return ImageFont.truetype(str(font_path), minimum_size)

    def _render_cover_page(self, metadata: Dict[str, Any]) -> io.BytesIO:
        """Flatten the approved template artwork and dynamic labels to one image.

        The template's original floating shapes are not portable across Word
        renderers. A single inline image preserves the exact visual hierarchy in
        desktop Word, Word Online, SharePoint previews and PDF converters.
        """
        template = self._cover_template_path()
        artwork: Optional[PILImage.Image] = None
        if template.exists():
            with zipfile.ZipFile(template, "r") as package:
                candidates = []
                for name in package.namelist():
                    if not name.startswith("word/media/"):
                        continue
                    try:
                        image = PILImage.open(io.BytesIO(package.read(name))).convert("RGBA")
                    except Exception:
                        continue
                    width, height = image.size
                    if height > width:
                        candidates.append((width * height, image))
                if candidates:
                    artwork = max(candidates, key=lambda item: item[0])[1]
        if artwork is None:
            fallback = Path(getattr(self.config, "COVER_PAGE_IMAGE", ""))
            if not fallback.exists():
                fallback = Path(self.config.ASSETS_DIR) / "coverpage.png"
            artwork = PILImage.open(fallback).convert("RGBA")

        canvas = artwork.copy()
        width, height = canvas.size
        logo_path = Path(self.config.ASSETS_DIR) / "ShellKode.png"
        if logo_path.exists():
            logo = PILImage.open(logo_path).convert("RGBA")
            bbox = logo.getbbox()
            if bbox:
                logo = logo.crop(bbox)
            target_width = round(width * 0.46)
            target_height = round(logo.height * target_width / logo.width)
            logo = logo.resize((target_width, target_height), PILImage.Resampling.LANCZOS)
            canvas.alpha_composite(logo, (round(width * 0.09), round(height * 0.035)))

        draw = ImageDraw.Draw(canvas)
        fonts_dir = Path(self.config.ASSETS_DIR) / "fonts"
        regular_path = fonts_dir / "DMSans-Regular.ttf"
        bold_path = fonts_dir / "DMSans-Bold.ttf"
        left = round(width * 0.095)
        maximum = round(width * 0.78)
        company = str(metadata.get("company_name") or "Client")
        project = str(metadata.get("project_title") or "Statement of Work")
        author = str(metadata.get("author_name") or "")
        organisation = str(metadata.get("author_org") or "ShellKode Pvt Ltd")
        if organisation.casefold() == "shellkode":
            organisation = "ShellKode Pvt Ltd"
        document_date = str(metadata.get("document_date") or "")

        company_font = self._fit_cover_text(draw, company, bold_path, round(height * 0.032), maximum)
        project_font = self._fit_cover_text(draw, project, regular_path, round(height * 0.020), maximum)
        detail_bold = ImageFont.truetype(str(bold_path), round(height * 0.014))
        detail = ImageFont.truetype(str(regular_path), round(height * 0.013))
        draw.text((left, round(height * 0.165)), company, font=company_font, fill="#7F00FF")
        draw.text((left, round(height * 0.225)), project, font=project_font, fill="#222222")
        lower_y = round(height * 0.685)
        draw.text((left, lower_y), organisation, font=detail_bold, fill="#FFFFFF")
        draw.text((left, lower_y + round(height * 0.038)), "Prepared by:", font=detail, fill="#FFFFFF")
        draw.text((left, lower_y + round(height * 0.068)), author, font=detail, fill="#FFFFFF")
        draw.text((left, lower_y + round(height * 0.098)), document_date, font=detail, fill="#FFFFFF")

        stream = io.BytesIO()
        canvas.convert("RGB").save(stream, format="PNG", optimize=True)
        stream.seek(0)
        return stream

    def _load_editable_cover(self, metadata: Dict[str, Any]) -> Document:
        """Load the approved Word cover and replace only its editable labels.

        The prior implementation rasterised the entire page. Besides making the
        labels uneditable, that page-height image could overflow Word Online's
        layout box and create a blank page before the TOC. Keeping the template's
        native paragraphs and anchored artwork preserves its geometry without a
        page-sized inline object.
        """
        template = self._cover_template_path()
        if not template.is_file():
            raise FileNotFoundError(f"Cover page template not found: {template}")
        document = Document(str(template))

        def replace_label(
            original: Any,
            value: str,
            size: float,
            *,
            bold: Optional[bool] = None,
            color: Optional[str] = None,
        ):
            accepted_labels = (original,) if isinstance(original, str) else tuple(original)
            for paragraph in document.paragraphs:
                if paragraph.text.strip() not in accepted_labels:
                    continue
                text_runs = [run for run in paragraph.runs if not run._r.xpath(".//w:drawing")]
                run = text_runs[0] if text_runs else paragraph.add_run()
                run.text = value
                for extra in text_runs[1:]:
                    extra.text = ""
                _set_font(run, FONT_NAME, size=size, bold=bold, color=color)
                return paragraph
            raise ValueError(
                "Cover template label not found; expected one of: "
                + ", ".join(str(label) for label in accepted_labels)
            )

        company = str(metadata.get("company_name") or "Client")
        project = str(metadata.get("project_title") or "Statement of Work")
        author = str(metadata.get("author_name") or "")
        organisation = str(metadata.get("author_org") or "ShellKode Pvt Ltd")
        if organisation.casefold() == "shellkode":
            organisation = "ShellKode Pvt Ltd"
        document_date = str(metadata.get("document_date") or "")

        # Keep only the approved template artwork. Flow paragraphs and framed
        # paragraphs are interpreted differently by Word and SharePoint. Use
        # modern DrawingML text boxes anchored to the physical page instead;
        # their w:t values remain directly editable in both editors.
        artwork_paragraph = next((p for p in document.paragraphs if p._p.xpath(".//w:drawing")), None)
        if artwork_paragraph is None:
            raise ValueError("Cover template contains no anchored artwork")
        for paragraph in list(document.paragraphs):
            if paragraph._p is not artwork_paragraph._p:
                paragraph._p.getparent().remove(paragraph._p)

        def add_textbox(name: str, text: str, x: float, y: float, width: float, height: float,
                        size: float, color: str, bold: bool = False) -> None:
            emu = lambda inches: str(round(inches * 914400))
            half_points = str(round(size * 2))
            bold_xml = "<w:b/><w:bCs/>" if bold else ""
            drawing = parse_xml(f"""
                <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                  <w:drawing>
                    <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"
                               relativeHeight="251660000" behindDoc="0" locked="0"
                               layoutInCell="1" allowOverlap="1">
                      <wp:simplePos x="0" y="0"/>
                      <wp:positionH relativeFrom="page"><wp:posOffset>{emu(x)}</wp:posOffset></wp:positionH>
                      <wp:positionV relativeFrom="page"><wp:posOffset>{emu(y)}</wp:posOffset></wp:positionV>
                      <wp:extent cx="{emu(width)}" cy="{emu(height)}"/>
                      <wp:effectExtent l="0" t="0" r="0" b="0"/>
                      <wp:wrapNone/>
                      <wp:docPr id="{100 + len(artwork_paragraph._p.xpath('.//wp:docPr'))}" name="{xml_escape(name)}"/>
                      <wp:cNvGraphicFramePr/>
                      <a:graphic>
                        <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                          <wps:wsp>
                            <wps:cNvSpPr txBox="1"/>
                            <wps:spPr>
                              <a:xfrm><a:off x="0" y="0"/><a:ext cx="{emu(width)}" cy="{emu(height)}"/></a:xfrm>
                              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                              <a:noFill/><a:ln><a:noFill/></a:ln>
                            </wps:spPr>
                            <wps:txbx>
                              <w:txbxContent>
                                <w:p>
                                  <w:pPr><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>
                                  <w:r>
                                    <w:rPr><w:rFonts w:ascii="DM Sans" w:hAnsi="DM Sans" w:eastAsia="DM Sans" w:cs="DM Sans"/>{bold_xml}<w:color w:val="{color}"/><w:sz w:val="{half_points}"/><w:szCs w:val="{half_points}"/></w:rPr>
                                    <w:t xml:space="preserve">{xml_escape(text)}</w:t>
                                  </w:r>
                                </w:p>
                              </w:txbxContent>
                            </wps:txbx>
                            <wps:bodyPr rot="0" spcFirstLastPara="1" vertOverflow="clip" horzOverflow="clip"
                                        vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0"
                                        numCol="1" anchor="t" anchorCtr="0"><a:spAutoFit/></wps:bodyPr>
                          </wps:wsp>
                        </a:graphicData>
                      </a:graphic>
                    </wp:anchor>
                  </w:drawing>
                </w:r>
            """)
            artwork_paragraph._p.append(drawing)

        add_textbox("Cover Company", company, 0.79, 3.42, 6.45, 0.62, 32, "7F00FF")
        add_textbox("Cover Project", project, 0.79, 4.18, 6.45, 0.40, 15, "222222")
        add_textbox("Cover Organisation", organisation, 0.79, 9.38, 4.75, 0.38, 14, "FFFFFF", True)
        add_textbox("Cover Prepared By", "Prepared by:", 0.79, 9.86, 4.75, 0.32, 12, "FFFFFF")
        add_textbox("Cover Author", author, 0.79, 10.22, 4.75, 0.32, 12, "FFFFFF")
        add_textbox("Cover Date", document_date, 0.79, 10.62, 4.75, 0.32, 12, "FFFFFF")
        return document

    def _add_content_section_after_cover(self):
        """Add the body section without introducing a blank spacer page.

        The retained cover is kept as its own Word section so the TOC and body
        retain their established page geometry, headers, footers and numbering.
        """
        assert self.doc is not None
        # The cover is A4 while the body is Letter. Word necessarily promotes a
        # continuous size-changing section to a new page; adding a manual break
        # as well therefore creates a blank page. Use one explicit next-page
        # section boundary and no additional break before the TOC.
        content_section = self.doc.add_section(WD_SECTION.NEW_PAGE)

        # python-docx represents a section break as a separate empty paragraph.
        # Word Online/SharePoint can lay that carrier out as its own page when
        # the cover already occupies the page. Attach the cover section
        # properties to the last cover paragraph and remove the empty carrier.
        # This is valid WordprocessingML and leaves no layout object that can
        # become a renderer-specific blank page.
        paragraphs = self.doc.paragraphs
        if len(paragraphs) >= 2:
            break_paragraph = paragraphs[-1]._p
            cover_paragraph = paragraphs[-2]._p
            break_ppr = break_paragraph.pPr
            sect_pr = break_ppr.sectPr if break_ppr is not None else None
            if sect_pr is not None and not break_paragraph.xpath(".//w:drawing"):
                cover_paragraph.get_or_add_pPr().append(sect_pr)
                break_paragraph.getparent().remove(break_paragraph)
        return content_section

    def _build_docx(self, output_path: Path, sections: Dict[str, Any], metadata: Dict[str, Any], mode: str) -> None:
        self.doc = self._load_editable_cover(metadata)
        self._figure_counter = 0
        self._table_counter = 0
        self._configure_styles()
        self._add_numbering()
        self.section_builder = SectionBuilder(self.doc, self.config)
        _mark_update_fields(self.doc)
        self.toc_entries = [self._resolve_title(entry, metadata) for entry in self.toc_entries]
        # The TOC is generated from the requested template before section text is
        # rendered.  If an upstream generator returns an empty optional section,
        # do not leave a PAGEREF pointing at a bookmark that cannot exist.
        self.toc_entries = [
            entry
            for entry in self.toc_entries
            if self._resolve_section_content(entry, sections, metadata).strip()
        ]

        content_section = self._add_content_section_after_cover()
        content_section.page_width = Inches(PAGE_WIDTH_IN)
        content_section.page_height = Inches(PAGE_HEIGHT_IN)
        content_section.top_margin = Inches(TOP_MARGIN_IN)
        content_section.bottom_margin = Inches(BOTTOM_MARGIN_IN)
        content_section.left_margin = content_section.right_margin = Inches(HORIZONTAL_MARGIN_IN)
        content_section.header_distance = Inches(0.28)
        content_section.footer_distance = Inches(0.25)
        content_section.different_first_page_header_footer = False
        content_section.header.is_linked_to_previous = False
        content_section.footer.is_linked_to_previous = False
        self._add_header_footer(content_section, metadata, mode)
        self._prepare_expanded_toc(sections, metadata)
        self._diagram_section_assignments = self._assign_diagram_sections(
            self._architecture_assets(sections), sections, metadata
        )
        self._inserted_diagram_ids: set[int] = set()

        # The next-page content section puts the TOC immediately after the
        # cover. Document Version Control and substantive sections follow it.
        self._add_toc()
        self.doc.add_page_break()
        for position, section_name in enumerate(self.toc_entries):
            self._build_section(section_name, sections, metadata, position)

        props = self.doc.core_properties
        props.author = metadata.get("author_name", "")
        props.last_modified_by = metadata.get("author_name", "")
        props.title = metadata.get("project_title", "")
        props.subject = f"{mode} Statement of Work for {metadata.get('company_name', '')}"
        props.keywords = "Statement of Work, SOW, AWS, ShellKode"
        props.comments = "Generated from source-grounded requirements; proposals and open clarifications are labelled."
        self.doc.save(str(output_path))
        _embed_dm_sans(output_path, self.config)
        try:
            _refresh_pageref_cached_results(output_path, self.config)
        except Exception as exc:
            # Pagination is an optional optimisation. The DOCX already holds
            # valid dynamic PAGEREF/PAGE fields, which Word and Word Online can
            # calculate when opened. Never block SOW delivery on an external
            # office-suite installation.
            print(f"⚠ TOC page-number cache not precomputed: {exc}")
        print(f"✓ DOCX saved: {output_path}")

    def _add_header_footer(self, section, metadata: Dict[str, Any], mode: str) -> None:
        header_title = metadata.get("document_header_title")
        if not header_title:
            mode_label = {
                "POC": "POC",
                "PROD": "Production",
                "POC_TO_PROD": "POC to Production",
            }.get(mode, mode)
            project = metadata.get("project_title", "Statement of Work")
            header_title = project if "sow" in project.casefold() else f"{project} – SOW ({mode_label})"
        logo_path = Path(self.config.ASSETS_DIR) / "ShellKode.png"
        org = self._get_short_company_name(metadata.get("author_org", "ShellKode"))
        year_match = re.search(r"\b(20\d{2})\b", str(metadata.get("document_date", "")))
        year = year_match.group(1) if year_match else str(datetime.now().year)

        def populate_header(header) -> None:
            trailing_paragraph = header.paragraphs[0]
            trailing_paragraph.paragraph_format.space_before = Pt(0)
            trailing_paragraph.paragraph_format.space_after = Pt(0)
            trailing_paragraph.paragraph_format.line_spacing = Pt(1)
            header_table = header.add_table(
                rows=1, cols=3, width=Inches(PAGE_WIDTH_IN - 2 * HORIZONTAL_MARGIN_IN)
            )
            header_table.alignment = WD_TABLE_ALIGNMENT.LEFT
            header_table.autofit = False
            header_widths = [7200, 1400, CONTENT_WIDTH_DXA - 8600]
            SectionBuilder._set_table_geometry(header_table, header_widths)
            self._remove_table_borders(header_table)
            header._element.remove(header_table._tbl)
            header._element.insert(0, header_table._tbl)
            left_cell, spacer_cell, right_cell = header_table.rows[0].cells
            for cell, width in zip(
                (left_cell, spacer_cell, right_cell), header_widths
            ):
                SectionBuilder._set_cell_width(cell, width)
                self._set_zero_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            left_paragraph = left_cell.paragraphs[0]
            spacer_paragraph = spacer_cell.paragraphs[0]
            right_paragraph = right_cell.paragraphs[0]
            left_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            spacer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            right_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for paragraph in (left_paragraph, spacer_paragraph, right_paragraph):
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
            left_paragraph.paragraph_format.left_indent = Inches(CHROME_INSET_IN)
            right_paragraph.paragraph_format.right_indent = Inches(CHROME_INSET_IN)
            title_run = left_paragraph.add_run(str(header_title))
            _set_font(title_run, size=6.5, bold=True, color=BLUE)
            if logo_path.exists():
                logo_run = right_paragraph.add_run()
                shape = logo_run.add_picture(str(logo_path), width=Inches(0.72))
                shape._inline.docPr.set("descr", "ShellKode")
            else:
                logo_text = right_paragraph.add_run(org)
                _set_font(logo_text, size=7, bold=True, color=BLUE)

        def populate_footer(footer) -> None:
            copyright_paragraph = footer.paragraphs[0]
            copyright_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            copyright_paragraph.paragraph_format.left_indent = Inches(CHROME_INSET_IN)
            copyright_paragraph.paragraph_format.space_before = Pt(2)
            copyright_paragraph.paragraph_format.space_after = Pt(0)
            copyright_paragraph.paragraph_format.line_spacing = 1.0
            self._set_paragraph_top_border(copyright_paragraph, BORDER)
            copyright_run = copyright_paragraph.add_run(
                f"Confidential Copyright © {org} {year}"
            )
            _set_font(copyright_run, size=5.75, color=MUTED)

            page_paragraph = footer.add_paragraph()
            page_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            page_paragraph.paragraph_format.space_before = Pt(0)
            page_paragraph.paragraph_format.space_after = Pt(0)
            page_paragraph.paragraph_format.line_spacing = 1.0
            page_label = page_paragraph.add_run("Page ")
            _set_font(page_label, size=5.75, color=MUTED)
            _add_field(page_paragraph, "PAGE", "1", size=5.75)
            of = page_paragraph.add_run(" of ")
            _set_font(of, size=5.75, color=MUTED)
            _add_field(page_paragraph, "NUMPAGES", "1", size=5.75)

        populate_header(section.header)
        populate_footer(section.footer)

    @staticmethod
    def _remove_table_borders(table) -> None:
        tblpr = table._tbl.tblPr
        borders = tblpr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblpr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = borders.find(qn(f"w:{edge}"))
            if element is None:
                element = OxmlElement(f"w:{edge}")
                borders.append(element)
            element.set(qn("w:val"), "nil")

    @staticmethod
    def _set_zero_cell_margins(cell) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        margins = tcpr.find(qn("w:tcMar"))
        if margins is None:
            margins = OxmlElement("w:tcMar")
            tcpr.append(margins)
        for side in ("top", "bottom", "start", "end"):
            element = margins.find(qn(f"w:{side}"))
            if element is None:
                element = OxmlElement(f"w:{side}")
                margins.append(element)
            element.set(qn("w:w"), "0")
            element.set(qn("w:type"), "dxa")

    def _prepare_expanded_toc(
        self, sections: Dict[str, Any], metadata: Dict[str, Any]
    ) -> None:
        """Build a TOC whose every displayed row has an independent page field."""
        self.expanded_toc_entries = []
        self.subheading_anchor_maps = {}
        for position, entry in enumerate(self.toc_entries, 1):
            top_anchor = _bookmark_name(entry, position)
            self.expanded_toc_entries.append((entry, top_anchor, 1))
            content = self._resolve_section_content(entry, sections, metadata)
            major_match = re.match(r"^\s*(\d+)[.)]?\s+", entry)
            major = major_match.group(1) if major_match else None
            heading_map: Dict[str, str] = {}
            for heading_index, (label, level) in enumerate(
                SectionBuilder.enumerate_headings(content, major), 1
            ):
                anchor = _bookmark_name(label, position * 1000 + heading_index)
                heading_map[label] = anchor
                self.expanded_toc_entries.append((label, anchor, min(level, 4)))
            self.subheading_anchor_maps[entry] = heading_map

    @staticmethod
    def _set_paragraph_top_border(paragraph, color: str) -> None:
        ppr = paragraph._element.get_or_add_pPr()
        borders = ppr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            ppr.append(borders)
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "4")
        top.set(qn("w:space"), "4")
        top.set(qn("w:color"), color)
        borders.append(top)

    @staticmethod
    def _set_paragraph_bottom_border(paragraph, color: str) -> None:
        ppr = paragraph._element.get_or_add_pPr()
        borders = ppr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            ppr.append(borders)
        bottom = borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), color)

    def _add_toc(self) -> None:
        assert self.doc is not None
        # The title must not itself be a Heading paragraph, otherwise a native
        # TOC refresh inserts "Table of Contents" into its own result.
        title = self.doc.add_paragraph()
        title.paragraph_format.keep_with_next = True
        title.paragraph_format.space_after = Pt(8)
        _set_font(title.add_run("Table of Contents"), size=20.0, bold=True, color=PURPLE)
        self._set_paragraph_bottom_border(title, PURPLE)
        entries = self.expanded_toc_entries or [
            (entry, _bookmark_name(entry, index), 1)
            for index, entry in enumerate(self.toc_entries, 1)
        ]
        if not entries:
            paragraph = self.doc.add_paragraph()
            _start_complex_field(paragraph, 'TOC \\o "1-4" \\h \\z \\u')
            paragraph.add_run("Update table of contents")
            _end_complex_field(paragraph)
            return
        entries_are_numbered = any(re.match(r"^\d+[.)]\s+", entry[0]) for entry in entries if entry[2] == 1)
        toc_paragraphs = []
        for index, (entry, anchor, level) in enumerate(entries, 1):
            paragraph = self.doc.add_paragraph()
            toc_paragraphs.append(paragraph)
            paragraph.paragraph_format.left_indent = Inches(
                0.02 if level == 1 else 0.22 + (level - 2) * 0.18
            )
            paragraph.paragraph_format.space_after = Pt(3 if level == 1 else 2)
            paragraph.paragraph_format.line_spacing = 1.15
            if index == 1:
                # The existing rows are the field's cached result, so generated
                # files display immediately. Word and Google Docs can replace
                # the complete result from Heading 1-4 with one update action.
                _start_complex_field(paragraph, 'TOC \\o "1-4" \\h \\z \\u')
            display = entry if entries_are_numbered or level > 1 else f"{index}. {entry}"
            _add_hyperlink(paragraph, display, anchor)
            toc_size = 11.0
            # Keep the right-aligned cached PAGEREF result inside the writable
            # page width. A hard-coded 7-inch stop sits too close to (and can
            # exceed after import) the right margin in Word/Google Docs.
            toc_right_edge = PAGE_WIDTH_IN - (2 * HORIZONTAL_MARGIN_IN) - 0.18
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Inches(toc_right_edge), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
            )
            tab = paragraph.add_run("\t")
            _set_font(tab, size=toc_size)
            _add_field(paragraph, f"PAGEREF {anchor} \\h", "", size=toc_size)
        _end_complex_field(toc_paragraphs[-1])

    def _resolve_content(self, section_name: str, sections: Dict[str, Any]) -> str:
        key = self._name_to_key(section_name)
        aliases = {
            "document_control": ("document_control_and_basis",),
            "document_version_control": ("document_control_and_basis",),
            "objective": ("project_overview", "project_overview_objectives"),
            "purpose_and_scope_of_this_deliverable": ("project_overview",),
            "current_state": ("current_state_and_business_context",),
            "executive_summary_and_project_overview": ("project_overview", "project_overview_objectives"),
            "detailed_scope_of_work": ("scope_of_work",),
            "scope_of_work": ("scope_of_work",),
            "detailed_production_scope_of_work": ("scope_of_work",),
            "architecture_overview": ("architecture_diagram", "architecture_integrations"),
            "solution_architecture_aws": ("architecture_diagram", "architecture_integrations"),
            "architecture_and_integrations": ("architecture_integrations", "architecture_diagram"),
            "assumptions_and_dependencies": ("assumptions",),
            "timeline_and_deliverables": ("timelines_and_deliverables",),
            "terms_and_conditions": ("terms_conditions",),
        }
        if key.endswith("_project_team_effort"):
            aliases[key] = ("shellkode_implementation_cost", "implementation_cost")
        for candidate in (key,) + aliases.get(key, ()):
            if sections.get(candidate):
                return str(sections[candidate])
        for candidate, content in sections.items():
            if candidate not in self.RESERVED_KEYS and content and self._is_similar(section_name, candidate):
                return str(content)
        return ""

    def _resolve_section_content(
        self,
        section_name: str,
        sections: Dict[str, Any],
        metadata: Dict[str, Any],
    ) -> str:
        """Resolve dynamic About headings as well as ordinary section aliases."""
        unnumbered_name = re.sub(
            r"^\s*\d+(?:\.\d+)*[.)]?\s*", "", section_name
        ).strip()
        author_about = f"About {self._get_short_company_name(metadata.get('author_org', 'ShellKode'))}"
        customer_about = f"About {metadata.get('company_name', 'Customer')}"
        if unnumbered_name.casefold() == author_about.casefold():
            return str(sections.get("about_shellkode", ""))
        if unnumbered_name.casefold() == customer_about.casefold():
            return str(
                sections.get("about_company")
                or sections.get("about_client")
                or ""
            )
        return self._resolve_content(section_name, sections)

    @staticmethod
    def _resolve_title(title: str, metadata: Dict[str, Any]) -> str:
        replacements = {
            "{AUTHOR_ORG_SHORT}": DocumentBuilder._get_short_company_name(metadata.get("author_org", "ShellKode")),
            "{COMPANY_NAME}": metadata.get("company_name", "Customer"),
            "{COMPANY_NAME_SHORT}": DocumentBuilder._get_short_company_name(metadata.get("company_name", "Customer")),
            "{PROJECT_TITLE}": metadata.get("project_title", "Project"),
        }
        resolved = title
        for placeholder, value in replacements.items():
            resolved = resolved.replace(placeholder, str(value))
        return resolved

    def _build_section(self, section_name: str, sections: Dict[str, Any], metadata: Dict[str, Any], position: int) -> None:
        assert self.doc is not None and self.section_builder is not None
        content = self._resolve_section_content(section_name, sections, metadata)
        if not content:
            print(f"⚠ No content for {section_name!r}; omitted from body")
            return
        heading = self.doc.add_paragraph(style="Heading 1")
        _set_font(heading.add_run(section_name), size=20.0, bold=True, color=PURPLE)
        self._set_paragraph_bottom_border(heading, PURPLE)
        self._add_bookmark_to_heading(heading, section_name, position + 1)
        assignments = getattr(self, "_diagram_section_assignments", None)
        if assignments is None:
            # Keep direct section-building calls used by integrations/tests
            # backward compatible; full document builds always pre-assign.
            architecture_assets = (
                self._architecture_assets(sections)
                if "architecture" in self._name_to_key(section_name) else []
            )
            self._inserted_diagram_ids = set()
        else:
            architecture_assets = [
                asset for asset in self._architecture_assets(sections)
                if assignments.get(id(asset)) == section_name
            ]

        def insert_for_heading(heading_text: str, _level: int):
            heading_key = self._name_to_key(heading_text)
            inserted = []
            for asset in architecture_assets:
                if id(asset) in self._inserted_diagram_ids:
                    continue
                placement_key = self._name_to_key(str(asset.get("placement_heading") or ""))
                type_key = self._name_to_key(str(asset.get("diagram_type") or ""))
                placement_tokens = set(placement_key.split("_")) - {"and", "the", "of"}
                heading_tokens = set(heading_key.split("_")) - {"and", "the", "of"}
                matches = bool(placement_tokens & heading_tokens) or (
                    type_key == "data_flow" and "flow" in heading_tokens
                ) or (
                    type_key == "integration_context" and "integration" in heading_tokens
                ) or (
                    type_key == "deployment_topology" and "deployment" in heading_tokens
                )
                if matches:
                    inserted.append(self._add_architecture_diagram(asset))
                    self._inserted_diagram_ids.add(id(asset))
            return [item for item in inserted if item is not None]
        major_match = re.match(r"^\s*(\d+)[.)]?\s+", section_name)
        major = major_match.group(1) if major_match else None
        self.section_builder.parse_content(
            content,
            heading_prefix=major,
            heading_anchor_map=self.subheading_anchor_maps.get(section_name, {}),
            bookmark_callback=self._add_bookmark_with_anchor,
            section_title=section_name,
            table_caption_callback=None,
            heading_callback=insert_for_heading,
        )
        for asset in architecture_assets:
            if id(asset) not in self._inserted_diagram_ids:
                self._add_architecture_diagram(asset)
                self._inserted_diagram_ids.add(id(asset))

    @staticmethod
    def _architecture_assets(sections: Dict[str, Any]) -> List[Dict[str, Any]]:
        assets = sections.get(ARCHITECTURE_ASSETS_KEY)
        if isinstance(assets, list):
            return [asset for asset in assets if isinstance(asset, dict)]
        legacy = sections.get(ARCHITECTURE_ASSET_KEY)
        return [legacy] if isinstance(legacy, dict) else []

    def _assign_diagram_sections(
        self,
        assets: List[Dict[str, Any]],
        sections: Dict[str, Any],
        metadata: Dict[str, Any],
    ) -> Dict[int, str]:
        """Assign each diagram to the closest substantive SOW section."""
        if not self.toc_entries:
            return {}
        stop_words = {"and", "the", "of", "to", "for", "a", "an", "diagram", "overview"}

        def tokens(value: Any) -> set[str]:
            return {
                token for token in re.findall(r"[a-z0-9]+", str(value or "").casefold())
                if len(token) > 2 and token not in stop_words
            }

        type_hints = {
            "data_flow": {"flow", "workflow", "process", "scope", "current", "state"},
            "integration_context": {"integration", "interface", "dependency", "dependencies", "system"},
            "deployment_topology": {"deployment", "implementation", "infrastructure", "architecture"},
        }
        assignments: Dict[int, str] = {}
        for asset in assets:
            diagram_type = str(asset.get("diagram_type") or "architecture_overview")
            if diagram_type == "architecture_overview":
                target = next(
                    (entry for entry in self.toc_entries if "solution_architecture" in self._name_to_key(entry)),
                    next((entry for entry in self.toc_entries if "architecture" in self._name_to_key(entry)), self.toc_entries[0]),
                )
                assignments[id(asset)] = target
                continue

            wanted = tokens(asset.get("placement_heading")) | tokens(asset.get("title"))
            hints = type_hints.get(diagram_type, set())
            scored = []
            for index, entry in enumerate(self.toc_entries):
                title_tokens = tokens(entry)
                content_tokens = tokens(self._resolve_section_content(entry, sections, metadata)[:5000])
                score = 8 * len(wanted & title_tokens)
                score += 3 * len(hints & title_tokens)
                score += len(wanted & content_tokens)
                score += len(hints & content_tokens)
                is_architecture = "architecture" in title_tokens
                scored.append((score, not is_architecture, -index, entry))
            assignments[id(asset)] = max(scored)[-1]
        return assignments

    def _add_table_caption(self, _table: Any, context: str):
        assert self.doc is not None
        self._table_counter += 1
        caption = self.doc.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption.add_run(f"Table {self._table_counter}: {context}")
        return caption

    def _add_architecture_diagram(self, asset: Any):
        """Insert a generated architecture image and durable edit link."""
        if not isinstance(asset, dict) or not asset.get("image_base64"):
            return None
        assert self.doc is not None
        try:
            image_bytes = base64.b64decode(asset["image_base64"], validate=True)
            with PILImage.open(io.BytesIO(image_bytes)) as image:
                pixel_width, pixel_height = image.size
            aspect = pixel_width / max(pixel_height, 1)
            width = min(6.95, 4.75 * aspect)
            height = width / max(aspect, 0.1)
            if height > 4.75:
                height = 4.75
                width = height * aspect

            explanation = str(
                asset.get("description")
                or (asset.get("spec") or {}).get("description")
                or f"This diagram depicts {asset.get('title') or 'the proposed workflow'}."
            ).strip()
            explanation_paragraph = self.doc.add_paragraph()
            explanation_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            explanation_paragraph.paragraph_format.space_before = Pt(7)
            explanation_paragraph.paragraph_format.space_after = Pt(5)
            explanation_paragraph.paragraph_format.line_spacing = 1.15
            label = explanation_paragraph.add_run(f"{asset.get('title') or 'Diagram'}: ")
            _set_font(label, size=11, bold=True, color=INK)
            explanation_run = explanation_paragraph.add_run(explanation)
            _set_font(explanation_run, size=11, color=INK)

            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(8)
            shape = paragraph.add_run().add_picture(
                io.BytesIO(image_bytes),
                width=Inches(width),
                height=Inches(height),
            )
            shape._inline.docPr.set(
                "descr",
                str(asset.get("alt_text") or "Proposed logical architecture diagram"),
            )

            edit_link = str(asset.get("edit_url") or "")
            if edit_link:
                edit_paragraph = self.doc.add_paragraph()
                edit_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                edit_paragraph.paragraph_format.space_before = Pt(2)
                edit_paragraph.paragraph_format.space_after = Pt(8)
                _add_external_hyperlink(
                    edit_paragraph, "Edit this diagram in draw.io", edit_link
                )
            return paragraph
        except Exception as exc:
            print(f"   ⚠ Architecture image could not be embedded; continuing without it: {exc}")
            return None

    def _add_bookmark_with_anchor(self, paragraph, anchor: str) -> None:
        self._bookmark_ids += 1
        bookmark_id = self._bookmark_ids
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), anchor)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def _add_bookmark_to_heading(self, heading, section_name: str, position: Optional[int] = None) -> None:
        self._bookmark_ids += 1
        bookmark_id = self._bookmark_ids
        anchor = _bookmark_name(section_name, position or bookmark_id)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), anchor)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        heading._p.insert(0, start)
        heading._p.append(end)


def add_toc_field(doc):
    """Backward-compatible dynamic TOC field helper."""
    paragraph = doc.add_paragraph()
    _add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u', "Update field to refresh table of contents")
    return paragraph
