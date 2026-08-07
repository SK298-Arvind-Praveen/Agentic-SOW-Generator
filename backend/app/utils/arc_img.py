import subprocess
from jinja2 import Template
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------------------------
# DATA
# ---------------------------
data = {
    "company_name": "ABC Corporation",
    "project_title": "AI Vehicle Inspection System",
    "author_org": "Kogta Financial",
    "author_name": "Abdul Rasith",
    "date": "2026-04-15",
    "base_cover": "C:/Users/DS/Pictures/SOW/assets/coverpage.png",
}

# ---------------------------
# SECTIONS (5 PAGES)
# ---------------------------
data["sections"] = [
    {"title": "Executive Summary"},
    {"title": "Scope of Work"},
    {"title": "System Architecture"},
    {"title": "Implementation Plan"},
    {"title": "Pricing and Commercials"},
]

# ---------------------------
# COVER GENERATION (FIXED)
# ---------------------------
def generate_cover(data):
    img = Image.open(data["base_cover"]).convert("RGBA")
    W, H = img.size

    draw = ImageDraw.Draw(img)

    font_path = "C:/Windows/Fonts/arial.ttf"

    # Scale factors (based on A4 reference from your PDF builder)
    scale_x = W / 595
    scale_y = H / 842

    def from_pdf(x, y):
        """Convert ReportLab coords → PIL coords"""
        return int(x * scale_x), int(H - (y * scale_y))

    # Fonts (scaled properly)
    title_font = ImageFont.truetype(font_path, int(39 * scale_y))
    subtitle_font = ImageFont.truetype(font_path, int(16 * scale_y))
    body_font = ImageFont.truetype(font_path, int(11 * scale_y))
    small_font = ImageFont.truetype(font_path, int(10 * scale_y))

    # ---------------------------
    # COVER TEXT (POSITION FIXED)
    # ---------------------------

    # Company name - positioned in upper left area
    draw.text(
        from_pdf(50, 720),
        data["company_name"],
        fill=(123, 63, 242),
        font=title_font
    )

    # Project title - positioned below company name with proper spacing
    draw.text(
        from_pdf(50, 670),
        data["project_title"],
        fill=(102, 102, 102),
        font=subtitle_font
    )

    # Bottom block - white text on dark section (lower left corner)
    draw.text(from_pdf(50, 120), data["author_org"], fill=(255, 255, 255), font=body_font)
    draw.text(from_pdf(50, 95), "Prepared by", fill=(255, 255, 255), font=small_font)
    draw.text(from_pdf(50, 70), data["author_name"], fill=(255, 255, 255), font=small_font)

    output = "cover_final.png"
    img.save(output)

    print("✅ Cover generated correctly")
    return output


cover_path = generate_cover(data)

# ---------------------------
# CREATE DOCX DIRECTLY
# ---------------------------
doc = Document()

# PAGE 1: COVER - Set margins to 0 for this section
section = doc.sections[0]
section.top_margin = Inches(0)
section.bottom_margin = Inches(0)
section.left_margin = Inches(0)
section.right_margin = Inches(0)

# Add cover image (full width)
paragraph = doc.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = paragraph.add_run()
run.add_picture(cover_path, width=Inches(8.27))  # A4 width

# PAGE 2: TOC - Add new section with normal margins
section = doc.add_section()
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

# TOC Title
toc_title = doc.add_heading('Table of Contents', level=1)
toc_title_format = toc_title.runs[0].font
toc_title_format.color.rgb = RGBColor(123, 63, 242)
toc_title_format.size = Pt(20)

# Add TOC field (Word will auto-generate page numbers)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_toc(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    
    return paragraph

add_toc(doc)

# CONTENT PAGES
for idx, sec in enumerate(data["sections"]):
    # Add new section for each content page (maintains margins)
    if idx > 0:
        section = doc.add_section()
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    else:
        # First content section continues from TOC
        doc.add_paragraph()  # Small spacing
    
    # Section heading
    heading = doc.add_heading(sec["title"], level=1)
    heading_format = heading.runs[0].font
    heading_format.color.rgb = RGBColor(123, 63, 242)
    heading_format.size = Pt(20)
    
    # Add bookmark to heading for TOC linking
    p = heading._element
    bookmark_id = sec["title"].replace(' ', '_').replace('&', 'and')
    bookmark_start = OxmlElement('w:bookmarkStart')
    bookmark_start.set(qn('w:id'), str(hash(bookmark_id) % 1000000))
    bookmark_start.set(qn('w:name'), f"_Toc{bookmark_id}")
    bookmark_end = OxmlElement('w:bookmarkEnd')
    bookmark_end.set(qn('w:id'), str(hash(bookmark_id) % 1000000))
    p.insert(0, bookmark_start)
    p.append(bookmark_end)
    
    # Section content
    for i in range(8):
        p = doc.add_paragraph(
            f'This section describes {sec["title"]} in detail. '
            f'It includes structured explanation, implementation insights, '
            f'technical considerations, and business alignment for the SOW.'
        )
        p_format = p.paragraph_format
        p_format.space_after = Pt(8)
        for run in p.runs:
            run.font.size = Pt(11)

# Save document
doc.save("sow.docx")

print("✅ DOCX generated with full-width cover and auto-TOC: sow.docx")
print("📝 Note: Right-click the TOC in Word and select 'Update Field' to populate page numbers")